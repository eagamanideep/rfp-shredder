"""Build a Full Shred from an ExtractionResult using rules only (no AI).

Rule-based output fills every tab with source-grounded rows found by pattern matching. Judgement columns (risk
levels, decision relevance, contradictions) are marked for human review; the AI builder in shred_ai.py fills them.
"""

from __future__ import annotations

import io
import re
import time
from collections.abc import Iterator
from contextlib import contextmanager
from dataclasses import dataclass, field
from datetime import date, datetime

from .extract import ExtractionResult, split_sentences
from .shred_schema import SHOW_STOPPER_AREAS, TAB_BY_NAME, TABS, TabSpec

MAX_ROWS = 40
REVIEW = "Rule-based keyword match; confirm against the cited text."


@dataclass
class Shred:
    client: str
    rfp_number: str
    title: str
    source_file: str
    run_date: date
    mode: str
    tabs: dict[str, list[dict[str, str]]] = field(default_factory=dict)
    banners: dict[str, str] = field(default_factory=dict)
    guide_notes: dict[str, str] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)
    # True when the user typed the value, so AI-detected values must not replace it.
    client_overridden: bool = False
    rfp_overridden: bool = False
    title_overridden: bool = False
    # (step, seconds) in the order the steps ran.
    timings: list[tuple[str, float]] = field(default_factory=list)

    @contextmanager
    def timed(self, step: str) -> Iterator[None]:
        start = time.perf_counter()
        try:
            yield
        finally:
            self.timings.append((step, time.perf_counter() - start))

    @property
    def title_suffix(self) -> str:
        rfp = " ".join(p for p in (f"RFP {self.rfp_number}" if self.rfp_number else "", self.title) if p)
        return " | ".join(p for p in (self.client, rfp) if p)

    def banner_for(self, spec: TabSpec) -> str:
        if spec.name == "Workbook Guide":
            return self.opportunity_line()
        return self.banners.get(spec.name) or spec.banner

    def opportunity_line(self) -> str:
        readiness = next(iter(self.tabs.get("Shred Readiness", [])), {}).get("status", "Not assessed")
        due = next((r for r in self.tabs.get("Key Dates", []) if r.get("date_type") == "Proposal Due Date"), None)
        deadline = (due.get("quick_turn") or due.get("date")) if due else "Not identified"
        parts = [
            self.client,
            f"RFP {self.rfp_number}" if self.rfp_number else "",
            self.title,
            f"Shred run: {self.run_date.isoformat()}",
            f"Overall readiness: {readiness}",
            f"Proposal deadline: {deadline}.",
        ]
        return "Opportunity: " + " | ".join(p for p in parts if p)


def assign_ids(tabs: dict[str, list[dict[str, str]]], names: list[str] | None = None) -> None:
    for name in names or list(tabs):
        prefix = TAB_BY_NAME[name].prefix
        for i, row in enumerate(tabs[name], start=1):
            row["id"] = f"{prefix}-{i:03d}"


def read_objection_matrix(data: bytes, max_rows: int = 300) -> str:
    """Flatten an objection-matrix workbook to pipe-separated text for screening."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines: list[str] = []
    for ws in wb.worksheets:
        lines.append(f"## Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(v).strip() for v in row if v is not None and str(v).strip()]
            if cells:
                lines.append(" | ".join(cells))
            if len(lines) >= max_rows:
                return "\n".join(lines)
    return "\n".join(lines)


# ---------------------------------------------------------------- source hits


@dataclass
class Hit:
    text: str
    page: int | None
    heading: str
    path: str

    @property
    def location(self) -> str:
        heading = self.heading if len(self.heading) <= 70 else self.heading[:67] + "…"
        return ", ".join(p for p in (f"p. {self.page}" if self.page else "", heading) if p)


class Context:
    def __init__(self, result: ExtractionResult, source_file: str, run_date: date):
        self.result = result
        self.source_file = source_file
        self.run_date = run_date
        self.hits = [
            Hit(sentence, p.page, s.heading, s.path)
            for s in result.sections
            for p in s.paragraphs
            for sentence in split_sentences(p.text)
        ]
        self.full_text = "\n".join(h.text for h in self.hits)
        self.requirements: list[dict[str, str]] = []

    def find(self, pattern: str, limit: int = MAX_ROWS, within: str | None = None) -> list[Hit]:
        rx = re.compile(pattern, re.I)
        scope = re.compile(within, re.I) if within else None
        seen: set[str] = set()
        out = []
        for h in self.hits:
            if rx.search(h.text) and (scope is None or scope.search(h.path)) and h.text not in seen:
                seen.add(h.text)
                out.append(h)
                if len(out) >= limit:
                    break
        return out

    def related(self, pattern: str, limit: int = 6) -> str:
        rx = re.compile(pattern, re.I)
        ids = [r["id"] for r in self.requirements if rx.search(r["requirement_text"])]
        return "; ".join(ids[:limit])

    def base(self, hit: Hit | None) -> dict[str, str]:
        return {"source_file": self.source_file, "source_location": hit.location if hit else ""}


# ---------------------------------------------------------------- opportunity metadata

_RFP_NUMBER_RE = re.compile(
    r"\b(?:RFP|RFQ|RFI|IFB|ITB|Solicitation|Bid)\s*(?:No\.?|Number|#|:)?\s*([A-Z]{0,8}[-\s]?[A-Z0-9]*\d[A-Z0-9\-_.]*)"
)
_CLIENT_RE = re.compile(r"\b(DEPARTMENT|SCHOOL|COUNTY|DISTRICT|CITY OF|UNIVERSITY|AUTHORITY|COMMISSION|BOARD)\b", re.I)
_SMALL_WORDS = {"of", "and", "the", "for", "in", "on", "at", "to"}


def _title_case(text: str) -> str:
    if text.upper() != text:
        return text
    words = text.lower().split()
    return " ".join(w if i and w in _SMALL_WORDS else w.capitalize() for i, w in enumerate(words))


def detect_opportunity(result: ExtractionResult, filename: str) -> tuple[str, str, str]:
    """Best-effort (client, RFP number, title) from the first pages; the UI lets users override them."""
    head = "\n".join(result.page_texts[:2])
    lines = [line.strip() for line in head.splitlines() if line.strip()]
    number = (m.group(1).strip(" .") if (m := _RFP_NUMBER_RE.search(head)) else "").strip()
    client = next((_title_case(line) for line in lines[:25] if _CLIENT_RE.search(line) and len(line) < 80), "")
    title = ""
    for i, line in enumerate(lines[:40]):
        if re.match(r"^REQUEST FOR (PROPOSALS?|QUOTES?|QUOTATIONS?|BIDS?|INFORMATION)", line, re.I):
            parts = []
            for nxt in lines[i + 1 : i + 5]:
                if _RFP_NUMBER_RE.search(nxt) or len(" ".join(parts + [nxt])) > 160:
                    break
                parts.append(nxt)
            title = " ".join(parts)
            break
    if not title:
        title = re.sub(r"[_\-]+", " ", filename.rsplit(".", 1)[0]).strip()
    return client, number, title


# ---------------------------------------------------------------- requirement classification

_TYPE_RULES = (
    ("Deadline", r"\b(due|deadline|no later than|late (bids|proposals)|received (by|until|no later))\b"),
    ("Q&A", r"\bquestions?\b|inquir"),
    ("Insurance", r"insurance|insured"),
    ("Notarization", r"notar|sworn"),
    ("Signature", r"\bsign(ed|ature|atory|s)?\b|\binitial"),
    ("Pricing", r"pric|\bcost|budget|invoice|payment|\bfees?\b"),
    ("Form", r"\bform\b|affidavit|W-9|certificate|appendix|attachment"),
    ("Submission", r"submi|envelope|\bcop(y|ies)\b|deliver|e-?mail|electronic|subject line"),
    ("Formatting", r"page limit|pages? or less|font|margin|typewritten|searchable|format"),
    ("Evaluation", r"evaluat|\bscor|\bpoints\b|weight"),
    ("Eligibility", r"licen[sc]|certif|eligib|regist|SPIN|E-?Rate"),
    ("Contract Term", r"terminat|assign|indemn|liabil|governed|warrant|breach|contract"),
    ("Reporting", r"report|meeting|status update|documentation|records"),
)
_WORKSTREAM = {
    "Deadline": "Submission",
    "Q&A": "Submission",
    "Submission": "Submission",
    "Formatting": "Submission",
    "Pricing": "Pricing",
    "Form": "Forms",
    "Signature": "Forms",
    "Notarization": "Forms",
    "Insurance": "Forms",
    "Contract Term": "Legal",
    "Eligibility": "Eligibility",
    "Evaluation": "Narrative",
}
_FORCE = {"Mandatory": "Mandatory", "Recommended": "Advisory", "Optional": "Conditional", "Commitment": "Informational"}


def _requirement_type(text: str) -> str:
    return next((name for name, rx in _TYPE_RULES if re.search(rx, text, re.I)), "Scope")


# ---------------------------------------------------------------- dates

_MONTH = r"(Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|June?|July?|Aug(?:ust)?|Sep(?:t(?:ember)?)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?)"
_DATE_RE = re.compile(
    rf"\b{_MONTH}\.?\s+(\d{{1,2}})(?:st|nd|rd|th)?,?\s+(\d{{4}})\b|\b(\d{{1,2}})/(\d{{1,2}})/(\d{{4}})\b"
)
_TIME_RE = re.compile(r"\b(\d{1,2}(?::\d{2})?\s*(?:a\.?m\.?|p\.?m\.?))", re.I)
_TZ_RE = re.compile(r"\b(EST|EDT|ET|CST|CDT|CT|MST|MDT|MT|PST|PDT|PT|Eastern|Central|Mountain|Pacific)\b")
_DATE_TYPES = (
    ("Q&A / Addendum Posting By", r"addend|answers?\b|responses to (inquiries|questions)"),
    ("Questions Due", r"question|inquir"),
    ("Bid Withdrawal Cutoff", r"withdraw"),
    ("Pre-Bid / Site Visit", r"pre-?bid|pre-?proposal|site visit|conference"),
    (
        "Proposal Due Date",
        r"proposals? due|due date|sealed (bids|proposals)|received (until|by|no later)|must be received|submission deadline|bids? due",
    ),
    ("Notification of Selection / Award", r"award|notification|selection"),
    ("RFP Release", r"releas|issued|advertis"),
    ("Contract Term / Completion", r"contract|term|through|completion|effective"),
)


def _parse_date(m: re.Match) -> date | None:
    try:
        if m.group(1):
            return datetime.strptime(f"{m.group(1)[:3]} {m.group(2)} {m.group(3)}", "%b %d %Y").date()
        return date(int(m.group(6)), int(m.group(4)), int(m.group(5)))
    except ValueError:
        return None


# ---------------------------------------------------------------- builder


def build_rule_shred(
    result: ExtractionResult,
    source_file: str,
    run_date: date,
    client: str = "",
    rfp_number: str = "",
    title: str = "",
    matrix_text: str | None = None,
) -> Shred:
    start = time.perf_counter()
    detected = detect_opportunity(result, source_file)
    shred = Shred(
        client=client or detected[0],
        rfp_number=rfp_number or detected[1],
        title=title or detected[2],
        source_file=source_file,
        run_date=run_date,
        mode="Rule-based",
        client_overridden=bool(client),
        rfp_overridden=bool(rfp_number),
        title_overridden=bool(title),
    )
    ctx = Context(result, source_file, run_date)
    tabs = shred.tabs

    tabs["Compliance Matrix"] = _compliance(ctx, shred)
    assign_ids(tabs, ["Compliance Matrix"])
    ctx.requirements = tabs["Compliance Matrix"]
    tabs["Key Dates"] = _key_dates(ctx)
    tabs["Scope Summary"] = _scope(ctx)
    tabs["Cost Breakdown"] = _cost(ctx)
    tabs["Client Objectives"] = _objectives(ctx)
    tabs["Evaluation Summary"] = _evaluation(ctx)
    tabs["Submission Mechanics"] = _submission(ctx)
    tabs["Formatting + Limits"] = _formatting(ctx)
    tabs["Accessibility 508 ADA"] = _presence_tab(
        ctx,
        r"section 508|\b508\b|\bADA\b|WCAG|accessib",
        lambda h: {
            "requirement": h.text,
            "standard": _first(r"Section 508|ADA|WCAG[\w .]*", h.text) or "Not stated",
            "applies_to": "Proposal and deliverables",
            "proposal_stage": "Unclear",
            "deliverable_stage": "Unclear",
        },
        {
            "requirement": "No Section 508, ADA, WCAG, accessible-document, or accessibility-conformance requirement "
            "was identified in the supplied document.",
            "standard": "Not stated",
            "applies_to": "Proposal and deliverables",
            "proposal_stage": "No",
            "deliverable_stage": "No",
            "evidence_artifact": "No accessibility artifact required by supplied document.",
            "notes": "Not Found; verify any addendum if issued.",
        },
    )
    tabs["Contract Vehicle Eligibility"] = _presence_tab(
        ctx,
        r"E-?Rate|\bSPIN\b|cooperative|contract vehicle|\bGSA\b|approved vendor|master agreement|piggyback",
        lambda h: {
            "eligibility_type": _first(r"E-?Rate|SPIN|cooperative|GSA|approved vendor|master agreement", h.text)
            or "Contract vehicle",
            "requirement": h.text,
            "applies_to": "Bidder",
            "mandatory": "Yes" if _mandatory(h.text) else "Unclear",
            "required_stage": _stage(h.text),
        },
        {
            "eligibility_type": "Contract vehicle / approved pool",
            "requirement": "No cooperative contract, master "
            "agreement, approved vendor pool, or task-order vehicle requirement was identified.",
            "applies_to": "Bidder",
            "mandatory": "No",
            "evidence_required": "None stated",
            "required_stage": "Not Applicable",
            "notes": "Not Found in supplied document.",
        },
    )
    tabs["Supplier Diversity Set-Aside"] = _presence_tab(
        ctx,
        r"\b(MBE|WBE|DBE|HUB|SBE|DVBE|minority|women-owned|disadvantaged|supplier diversity|set-aside)\b",
        lambda h: {
            "program_requirement": h.text,
            "type": _first(r"MBE|WBE|DBE|HUB|SBE|DVBE|set-aside", h.text) or "Diversity",
            "mandatory_goal": "Mandatory" if _mandatory(h.text) else "Goal",
            "percentage_threshold": _first(r"\d+(\.\d+)?\s*%", h.text) or "Not stated",
        },
        {
            "program_requirement": "Supplier diversity / set-aside",
            "type": "No supplier-diversity set-aside, "
            "MBE/WBE/DBE goal, participation percentage, or certification requirement was identified.",
            "mandatory_goal": "Not Found",
            "percentage_threshold": "Not stated",
            "certification_evidence": "None stated",
            "subcontracting_requirement": "None stated",
            "notes": "Not Found; verify any addendum.",
        },
    )
    tabs["Work Samples + References"] = _references(ctx)
    tabs["Forms + Attachments"] = _forms(ctx)
    tabs["Proposal Outline"] = _outline(ctx, tabs["Forms + Attachments"])
    tabs["Contract Objections"] = _objections(ctx, matrix_text)
    tabs["Source Inventory"] = _sources(ctx, matrix_text)
    tabs["Amendment + Delta Log"] = _amendments(ctx)
    tabs["Show Stoppers"] = _show_stoppers(ctx, tabs)
    tabs["Shred Readiness"] = _readiness(ctx, tabs, matrix_text)
    tabs["Q&A + Issue Log"] = _issues(ctx, tabs)
    for name, rows in tabs.items():  # absences stay visible rather than leaving a tab blank
        if not rows:
            first = TAB_BY_NAME[name].keys[0]
            rows.append(
                {
                    first: "Not Found",
                    "source_file": source_file,
                    "source_location": "Entire document reviewed",
                    "notes": f"No {name.lower()} content was identified by rule-based extraction; verify manually.",
                }
            )
    assign_ids(tabs, [t.name for t in TABS if t.name in tabs and t.name != "Compliance Matrix"])
    _finish_notes(shred)
    shred.warnings.append(
        "Rule-based shred: risk levels, decision relevance and contradictions are not assessed. "
        "Add an OpenAI API key in the sidebar for the full analysis."
    )
    shred.timings.append(("Rule-based shred (all tabs)", time.perf_counter() - start))
    return shred


def _finish_notes(shred: Shred) -> None:
    """Fill the guide notes and banners that depend on the finished tabs."""
    readiness = next(iter(shred.tabs.get("Shred Readiness", [])), {})
    if readiness:
        shred.guide_notes["Shred Readiness"] = f"Overall status is {readiness.get('status')}: " + readiness.get(
            "issue_limitation", ""
        )
        shred.banners["Shred Readiness"] = shred.guide_notes["Shred Readiness"]
    due = next((r for r in shred.tabs.get("Key Dates", []) if r.get("date_type") == "Proposal Due Date"), None)
    if due:
        shred.guide_notes["Key Dates"] = (
            f"Quick Turn = 7 or fewer calendar days. Proposal due date is {due.get('quick_turn') or 'not assessed'} "
            "on the shred run date."
        )


def _first(pattern: str, text: str) -> str:
    m = re.search(pattern, text, re.I)
    return m.group(0).strip() if m else ""


def _mandatory(text: str) -> bool:
    return bool(re.search(r"\b(must|shall|required|mandatory)\b", text, re.I))


def _stage(text: str) -> str:
    if re.search(
        r"award|successful (bidder|vendor|contractor|offeror)|prior to commencement|after (the )?contract", text, re.I
    ):
        return "Post-Award"
    return "Proposal Submission"


def _compliance(ctx: Context, shred: Shred) -> list[dict[str, str]]:
    rows = []
    for r in ctx.result.requirements:
        kind = _requirement_type(r.text)
        force = _FORCE.get(r.type, "Informational")
        if re.match(r"^\W*(if|where|when|unless|should a)\b", r.text, re.I) and force == "Mandatory":
            force = "Conditional"
        heading = r.section.split(" > ")[-1]
        rows.append(
            {
                "requirement_force": force,
                "requirement_type": kind,
                "requirement_text": r.text,
                "response_artifact_needed": "",
                "proposal_package_home": "",
                "workstream": _WORKSTREAM.get(kind, "Other"),
                "evaluation_link": "",
                "source_file": ctx.source_file,
                "source_location": ", ".join(p for p in (f"p. {r.page}" if r.page else "", heading) if p),
                "amendment_version": f"Baseline {shred.rfp_number}".strip(),
                "related_ids": "",
                "notes": f"Keyword: {r.keyword}",
            }
        )
    return rows


def _key_dates(ctx: Context) -> list[dict[str, str]]:
    rows = [
        {
            "date_type": "Shred Run Date",
            "date": ctx.run_date.strftime("%b %d, %Y"),
            "time_zone": "Not Applicable",
            "mandatory": "No",
            "calendar_days_from_shred_run": "0",
            "source_file": "Internal run control",
            "source_location": "Shred run date",
            "notes": "Baseline for calendar-day calculations.",
        }
    ]
    candidates = [(h.text, h.location) for h in ctx.hits]
    for table in ctx.result.tables:  # schedule tables keep their dates out of the body text
        for row in table.rows:
            candidates.append((" ".join(c for c in row if c), f"p. {table.page}, table" if table.page else "table"))
    seen: set[tuple[str, date]] = set()
    for text, location in candidates:
        for m in _DATE_RE.finditer(text):
            when = _parse_date(m)
            if not when:
                continue
            kind = next((name for name, rx in _DATE_TYPES if re.search(rx, text, re.I)), "Other Date")
            # Unlabelled dates far from the run date are form versions, statute dates and the like.
            if (kind, when) in seen or (kind == "Other Date" and abs((when - ctx.run_date).days) > 365):
                continue
            seen.add((kind, when))
            days = (when - ctx.run_date).days
            quick = ""
            if kind == "Proposal Due Date":
                quick = "Past Due" if days < 0 else "Yes" if days <= 7 else "No"
            mandatory = (
                "Yes"
                if kind in ("Proposal Due Date", "Questions Due", "Bid Withdrawal Cutoff")
                else ("Yes" if re.search(r"mandatory", text, re.I) else "Conditional" if "Addendum" in kind else "No")
            )
            rows.append(
                {
                    "date_type": kind,
                    "date": when.strftime("%b %d, %Y"),
                    "time": (t.group(1) if (t := _TIME_RE.search(text)) else ""),
                    "time_zone": (z.group(1) if (z := _TZ_RE.search(text)) else "Not stated"),
                    "mandatory": mandatory,
                    "calendar_days_from_shred_run": str(days),
                    "quick_turn": quick,
                    "source_file": ctx.source_file,
                    "source_location": location,
                    "notes": text[:300],
                }
            )
            if len(rows) > MAX_ROWS:
                return rows
    return rows


def _quantity_tables(ctx: Context):
    """Yield (page, description, quantity, unit) for tables that have a quantity column."""
    for table in ctx.result.tables:
        header = [c.lower() for c in table.rows[0]]
        qty = next((i for i, c in enumerate(header) if re.search(r"\bqty\b|quantity", c)), None)
        if qty is None:
            continue
        desc = next((i for i, c in enumerate(header) if re.search(r"desc|item|part|product|model", c) and i != qty), 0)
        unit = next((i for i, c in enumerate(header) if re.search(r"\bunit\b|\buom\b", c)), None)
        for row in table.rows[1:]:
            if row[desc] or row[qty]:
                yield (
                    table.page,
                    " — ".join(c for i, c in enumerate(row) if c and i not in (qty, unit)),
                    row[qty],
                    (row[unit] if unit is not None else "Each"),
                )


def _scope(ctx: Context) -> list[dict[str, str]]:
    rows = []
    for page, description, qty, unit in _quantity_tables(ctx):
        rows.append(
            {
                "scope_area": "Line item",
                "requirement_description": description,
                "quantity": qty,
                "unit": unit,
                "mandatory": "Yes",
                "source_file": ctx.source_file,
                "source_location": f"p. {page}, table" if page else "table",
            }
        )
    for h in ctx.find(
        r"\b(must|shall|will|required)\b", within=r"scope|statement of (work|purpose)|services|deliverable|requirement"
    ):
        if len(rows) >= MAX_ROWS:
            break
        rows.append(
            {
                "scope_area": h.heading,
                "requirement_description": h.text,
                "mandatory": "Yes" if _mandatory(h.text) else "Unclear",
                "cabling_installation": "Yes" if re.search(r"install|cabl", h.text, re.I) else "",
                **ctx.base(h),
            }
        )
    return rows


def _cost(ctx: Context) -> list[dict[str, str]]:
    rows = []
    for page, description, qty, unit in _quantity_tables(ctx):
        rows.append(
            {
                "cost_pricing_element": description,
                "quantity": qty,
                "unit": unit,
                "pricing_basis": "Unit price + total price",
                "budget_ceiling_nte": "Not Found",
                "separate_cost_file": "Unclear",
                "source_file": ctx.source_file,
                "source_location": f"p. {page}, table" if page else "table",
            }
        )
    budget = ctx.find(r"\bbudget\b|not[- ]to[- ]exceed|\bceiling\b|price limitation|maximum (amount|price)", limit=3)
    separate = bool(ctx.find(r"(price|cost) proposal .*(separate|separately)|separately .*(price|cost)", limit=1))
    for h in ctx.find(r"pric|\bcost|invoice|payment|\bfees?\b|reimburs|travel", limit=25):
        rows.append(
            {
                "cost_pricing_element": h.heading,
                "pricing_basis": h.text,
                "budget_ceiling_nte": "Found" if budget else "Not Found",
                "separate_cost_file": "Yes" if separate else "Unclear",
                "reimbursables_travel": h.text if re.search(r"reimburs|travel", h.text, re.I) else "",
                **ctx.base(h),
            }
        )
    if not budget:
        rows.append(
            {
                "cost_pricing_element": "Budget / ceiling / NTE",
                "unit": "Not stated",
                "pricing_basis": "No published affordability cap",
                "budget_ceiling_nte": "Not Found",
                "separate_cost_file": "Yes" if separate else "No",
                "source_file": ctx.source_file,
                "notes": "No budget, ceiling, or not-to-exceed amount was identified.",
            }
        )
    return rows


def _objectives(ctx: Context) -> list[dict[str, str]]:
    return [
        {
            "objective_pressure": h.text,
            "beneficiary_stakeholder": "",
            "confidence": "Moderate",
            "notes": f"Found under '{h.heading}'.",
            **ctx.base(h),
        }
        for h in ctx.find(r"\b(goal|objective|purpose|intend|seeks?|desired|outcome|improve)\b", limit=15)
    ]


def _evaluation(ctx: Context) -> list[dict[str, str]]:
    rows = []
    for table in ctx.result.tables:
        header = " ".join(table.rows[0]).lower()
        if not re.search(r"point|weight|score|%|criteria", header):
            continue
        for row in table.rows[1:]:
            weight = next((c for c in reversed(row) if re.fullmatch(r"\d{1,3}\s*(%|points?)?", c.strip(), re.I)), "")
            label = next((c for c in row if c and c != weight), "")
            if weight and label and not re.search(r"total|maximum", label, re.I):
                rows.append(
                    {
                        "criterion_factor": label,
                        "weight_points": weight,
                        "related_requirement_ids": ctx.related(re.escape(label[:30])),
                        "source_file": ctx.source_file,
                        "source_location": f"p. {table.page}, table" if table.page else "table",
                    }
                )
    for h in ctx.find(r"\b\d{1,3}\s*(%|percent|points)\b|weighted", within=r"evaluat|criteria|scor|selection|award"):
        rows.append(
            {
                "criterion_factor": h.heading,
                "weight_points": _first(r"\d{1,3}\s*(%|percent|points)", h.text),
                "what_evaluator_will_score": h.text,
                **ctx.base(h),
            }
        )
    return rows[:MAX_ROWS]


_SUBMISSION_CONTROLS = (
    ("Deadline receipt", r"received|\bdue\b|deadline|\blate\b"),
    ("Copies / originals", r"\bcop(y|ies)\b|original"),
    ("Physical delivery", r"sealed|envelope|mail|hand[- ]deliver"),
    ("Electronic submission", r"e-?mail|electronic|portal|upload|\bPDF\b|\bMB\b"),
    ("Labeling / subject line", r"label|subject line|marked|face of"),
    ("Signature control", r"\bsign|initial"),
    ("Notarization", r"notar|sworn"),
    ("Package order / content", r"order|organized|index|redact"),
)


def _submission(ctx: Context) -> list[dict[str, str]]:
    rows = []
    for h in ctx.find(r"submi|deliver|envelope|\bcop(y|ies)\b|e-?mail|portal|sealed|received"):
        control = next((name for name, rx in _SUBMISSION_CONTROLS if re.search(rx, h.text, re.I)), None)
        if not control:
            continue
        method = (
            "Electronic"
            if re.search(r"e-?mail|electronic|portal|upload|PDF", h.text, re.I)
            else ("Physical package" if re.search(r"sealed|envelope|mail|deliver", h.text, re.I) else "")
        )
        rows.append(
            {
                "control": control,
                "requirement": h.text,
                "required_stage": _stage(h.text),
                "submission_method": method,
                "copies_originals": _first(r"(\w+ \(\d+\)|\d+) (electronic )?(original|cop(y|ies))[^.;,]*", h.text),
                "signature_type": _first(r"blue ink|wet signature|notar\w*|electronic signature", h.text),
                "portal_delivery_details": _first(r"[\w.+-]+@[\w-]+\.[\w.]+(?:.*?@[\w-]+\.[\w.]+)?", h.text),
                "related_requirement_ids": ctx.related(re.escape(h.text[:40])),
                **ctx.base(h),
            }
        )
    return rows


_FORMAT_RULES = (
    ("Page limit", r"page limit|pages? or less|no more than \d+ pages|\(\d+\) pages"),
    ("Font / spacing", r"\bfont\b|point size|single-spaced|double-spaced"),
    ("Margins", r"margin"),
    ("File format", r"searchable|\bPDF\b|\.docx|\bWord format"),
    ("File size", r"\d+ ?MB|file size"),
    ("Completion method", r"typewritten|blue (pen|ink)"),
    ("Page order / indexing", r"numerical order|clearly indexed|organized in the following order"),
)


def _formatting(ctx: Context) -> list[dict[str, str]]:
    rows = []
    for h in ctx.hits:
        rule = next((name for name, rx in _FORMAT_RULES if re.search(rx, h.text, re.I)), None)
        if rule and len(rows) < MAX_ROWS:
            rows.append(
                {
                    "applies_to": "Proposal",
                    "rule_type": rule,
                    "requirement": h.text,
                    "limit_value": _first(
                        r"\d+\s*(pages?|MB|point)|five \(5\) pages|blue (pen|ink)|searchable PDF", h.text
                    ),
                    "consequence": h.text
                    if re.search(r"reject|disqualif|invalid|not be (scored|considered)", h.text, re.I)
                    else "",
                    "related_requirement_ids": ctx.related(re.escape(h.text[:40])),
                    **ctx.base(h),
                }
            )
    for rule, label in (("Page limit", "No proposal page limit"), ("Font / spacing", "No font or spacing rule")):
        if not any(r["rule_type"] == rule for r in rows):
            rows.append(
                {
                    "applies_to": "Proposal",
                    "rule_type": rule,
                    "requirement": f"{label} was identified in the supplied document.",
                    "limit_value": "Not Found",
                    "consequence": "No stated consequence.",
                    "source_file": ctx.source_file,
                    "notes": "Visible Not Found control; verify any addendum.",
                }
            )
    return rows


def _presence_tab(ctx: Context, pattern: str, make, not_found: dict[str, str]) -> list[dict[str, str]]:
    hits = ctx.find(pattern, limit=15)
    if not hits:
        return [{**not_found, "source_file": ctx.source_file, "source_location": "Entire document reviewed"}]
    return [{**make(h), "related_requirement_ids": ctx.related(pattern), "notes": REVIEW, **ctx.base(h)} for h in hits]


def _references(ctx: Context) -> list[dict[str, str]]:
    rows = []
    pattern = r"\breferences?\b|similar (goods|services|projects?|contracts?|scope)|examples? (of|that)|case stud|work sample|past performance|prior experience"
    for h in ctx.find(pattern, limit=20):
        kind = (
            "Reference"
            if re.search(r"reference", h.text, re.I)
            else ("Work sample" if re.search(r"sample|example|case stud", h.text, re.I) else "Similar-project evidence")
        )
        rows.append(
            {
                "item_type": kind,
                "requirement": h.text,
                "quantity": _first(
                    r"\b(one|two|three|four|five|\d+)\b(\s*\(\d+\))?(?=\s+(current|recent|references?|examples?|clients?|customers?))",
                    h.text,
                )
                or "Not stated",
                "recency_similarity": _first(r"current or recent|recent|similar[\w ]*", h.text) or "Not stated",
                "related_requirement_ids": ctx.related(pattern),
                **ctx.base(h),
            }
        )
    if not rows:
        rows.append(
            {
                "item_type": "Work samples",
                "requirement": "No reference, work-sample or similar-project requirement was identified.",
                "quantity": "Not stated",
                "source_file": ctx.source_file,
                "notes": "Not Found; verify any addendum.",
            }
        )
    return rows


_FORM_NAME_RE = re.compile(
    r"((?:[A-Z][\w\-/&'’.]*\s+){0,5}(?:Affidavit|Form|Certificate|Declaration|Application|Agreement)\b(?:\s*\([^)]{1,40}\))?"
    r"|\bW-9\b|\b(?:Appendix|Attachment|Exhibit|Form)\s+[A-Z0-9][\w\-]*)"
)


def _forms(ctx: Context) -> list[dict[str, str]]:
    rows: dict[str, dict[str, str]] = {}
    for h in ctx.find(
        r"affidavit|W-9|\bform\b|certificate|declaration|attachment|appendix|exhibit|application", limit=200
    ):
        for m in _FORM_NAME_RE.finditer(h.text):
            name = re.sub(r"^(The|A|An|This|Each|Any|Completed?|Signed)\s+", "", m.group(1).strip(), flags=re.I)
            key = name.lower()
            generic = key in {"form", "agreement", "application", "certificate", "declaration", "affidavit"}
            if generic or len(name) < 3 or key in rows or len(rows) >= MAX_ROWS:
                continue
            sig = (
                "Signature + notarization"
                if re.search(r"notar|sworn", h.text, re.I)
                else ("Signature" if re.search(r"\bsign", h.text, re.I) else "")
            )
            rows[key] = {
                "form_attachment": name,
                "required": "Yes" if _mandatory(h.text) else "Unclear",
                "required_stage": _stage(h.text),
                "signature_notary": sig,
                "who_completes": "Bidder",
                "separate_artifact": "Unclear",
                "file_provided": "Unclear",
                "related_requirement_ids": ctx.related(re.escape(name)),
                "notes": h.text[:250],
                **ctx.base(h),
            }
    return list(rows.values())


def _outline(ctx: Context, forms: list[dict[str, str]]) -> list[dict[str, str]]:
    rows = []
    marker = re.compile(
        r"proposal (content|format|organization|requirements|response)|submission requirements|response format", re.I
    )
    for parent in ctx.result.sections:
        if not marker.search(parent.heading):
            continue
        for s in ctx.result.sections:
            if s.path.startswith(parent.path + " > "):
                clean = re.sub(r"^([IVX]+|[A-Z]|\d+(\.\d+)*)[.)]\s+", "", s.heading)
                rows.append(
                    {
                        "toc_level": str(s.level - parent.level),
                        "structural_role": "Required Response Item",
                        "rfp_heading_item": s.heading,
                        "clean_proposal_heading": clean,
                        "requirement_ids": ctx.related(re.escape(clean[:30])),
                        "required": "Yes",
                        "narrative_artifact": "Artifact"
                        if re.search(r"budget|price|cost|form|cover|page", clean, re.I)
                        else "Narrative",
                        "source_file": ctx.source_file,
                        "source_location": f"p. {s.page}, {parent.heading}" if s.page else parent.heading,
                    }
                )
    for form in forms:
        rows.append(
            {
                "toc_level": "Forms",
                "structural_role": "Form",
                "rfp_heading_item": form["form_attachment"],
                "clean_proposal_heading": form["form_attachment"],
                "requirement_ids": form["related_requirement_ids"],
                "required": form["required"],
                "narrative_artifact": "Artifact",
                "source_file": ctx.source_file,
                "source_location": form["source_location"],
            }
        )
    return rows[:MAX_ROWS]


_CT_CATEGORIES = (
    (
        "Limitation of liability",
        r"limitation of liability|liability (shall be |is )?limited|consequential damages|not to exceed the total fees",
    ),
    ("Indemnification", r"indemnif|hold harmless"),
    ("Assignment", r"\bassign(ment|ed)?\b|delegat|change of control"),
    (
        "Payment terms",
        r"payment (shall|will) be made|\bnet \d+|withhold payment|payments? (shall be |are |is )?contingent|upon .*acceptance of each deliverable",
    ),
    ("Termination for convenience", r"terminat\w*[^.]*(convenience|for any reason)"),
    ("Governing law / venue", r"governed[^.]*laws|exclusive jurisdiction|venue"),
    (
        "Intellectual property / data ownership",
        r"shall be the property of|intellectual property|work product|ownership",
    ),
    ("Set-off / liquidated damages", r"liquidated|set off|offset"),
    ("Audit / records access", r"access to[^.]*(books|records)|\baudit\b"),
)


def _objections(ctx: Context, matrix_text: str | None) -> list[dict[str, str]]:
    acceptance = bool(
        ctx.find(
            r"by submitting|submission (of a proposal )?(constitutes|=) acceptance|agrees? to (all|the) terms|understands, agrees",
            limit=1,
        )
    )
    basis = (
        "Matrix supplied; rule-based mode does not apply matrix thresholds - use AI mode for matrix screening."
        if matrix_text
        else "Not screened against an internal objection matrix (none supplied)."
    )
    rows = []
    for category, pattern in _CT_CATEGORIES:
        hits = ctx.find(pattern, limit=1)
        if not hits:
            continue
        h = hits[0]
        rows.append(
            {
                "section_clause": h.heading,
                "clause_summary": "",
                "triggering_language_short_quote": h.text[:300],
                "trigger_category": category,
                "matrix_basis_threshold": basis,
                "client_restriction": "Submission = Acceptance" if acceptance else "",
                "matrix_escalation": "Unclear",
                "proposal_risk_level": "Watch",
                "required_internal_action": "Contracts Review",
                "related_requirement_ids": ctx.related(pattern),
                "confidence": "Low",
                "notes": REVIEW,
                **ctx.base(h),
            }
        )
    if not ctx.find(_CT_CATEGORIES[0][1], limit=1):
        rows.insert(
            0,
            {
                "section_clause": "Terms and conditions — absence of limitation-of-liability provision",
                "clause_summary": "No limitation-of-liability cap or damages limitation was identified.",
                "triggering_language_short_quote": "[No limitation-of-liability section identified]",
                "trigger_category": "Limitation of liability",
                "matrix_basis_threshold": basis,
                "client_restriction": "Submission = Acceptance" if acceptance else "",
                "matrix_escalation": "Unclear",
                "proposal_risk_level": "Watch",
                "required_internal_action": "Contracts Review",
                "source_file": ctx.source_file,
                "source_location": "Entire document reviewed",
                "confidence": "Low",
                "notes": REVIEW,
            },
        )
    return rows


def _sources(ctx: Context, matrix_text: str | None) -> list[dict[str, str]]:
    pages = len(ctx.result.page_texts)
    rows = [
        {
            "file_source_name": ctx.source_file,
            "source_type": "RFP / procurement package",
            "procurement_authority": "Yes",
            "reviewed": "Yes",
            "readability_access": f"Readable; {pages} pages" if ctx.result.file_type == "pdf" else "Readable",
            "superseded": "No",
            "use_scope": "Baseline procurement authority",
            "notes": "Uploaded source for this shred.",
        }
    ]
    if ctx.find(r"addend", limit=1):
        rows.append(
            {
                "file_source_name": "Addendum / Q&A release (if issued)",
                "source_type": "Referenced procurement update",
                "procurement_authority": "Yes",
                "reviewed": "No",
                "readability_access": "Not supplied; existence not confirmed",
                "superseded": "Unclear",
                "use_scope": "Could amend any procurement requirement",
                "notes": "Affected conclusions remain subject to update if an addendum exists.",
            }
        )
    referenced = sorted(
        {m.group(0) for h in ctx.hits for m in re.finditer(r"\b(Appendix|Attachment|Exhibit)\s+[A-Z0-9]+\b", h.text)}
    )
    for name in referenced[:10]:
        rows.append(
            {
                "file_source_name": name,
                "source_type": "Referenced attachment",
                "procurement_authority": "Yes",
                "reviewed": "Unclear",
                "readability_access": "Check whether included in the supplied file",
                "superseded": "Unclear",
                "use_scope": "Referenced in document",
            }
        )
    if matrix_text:
        rows.append(
            {
                "file_source_name": "Contracts Proposal Objection Matrix (uploaded)",
                "source_type": "Internal Contracts Proposal Objection Matrix",
                "procurement_authority": "No",
                "reviewed": "Yes",
                "readability_access": "Readable",
                "superseded": "No",
                "use_scope": "Sole internal authority for contract-objection trigger logic",
                "notes": "Not procurement authority; used only for internal contract-trigger screening.",
            }
        )
    return rows


def _amendments(ctx: Context) -> list[dict[str, str]]:
    rows = [
        {
            "change_source": ctx.source_file,
            "change_type": "Baseline established",
            "previous_rule_value": "Not Applicable",
            "new_rule_value": "Baseline captured from supplied source",
            "affected_ids": "All baseline IDs",
            "affected_tabs_workstreams": "All tabs",
            "action_required": "Preserve current IDs if a later addendum or reissue is provided",
            "notes": "No amendment/addendum file was supplied in this run.",
        }
    ]
    if hits := ctx.find(r"addend", limit=1):
        rows.append(
            {
                "change_source": "Potential Q&A / addendum release",
                "change_type": "Source gap; no delta confirmed",
                "previous_rule_value": "Baseline only",
                "new_rule_value": "No addendum text available to compare",
                "affected_ids": "All IDs potentially affected",
                "affected_tabs_workstreams": "All tabs / workstreams",
                "action_required": "Obtain or confirm absence of addenda; refresh affected tabs while preserving IDs",
                "source_location": hits[0].location,
                "notes": hits[0].text[:250],
            }
        )
    return rows


_SS_PATTERNS = {
    "Mandatory pre-bid / site visit / LOI": r"pre-?bid|pre-?proposal conference|site visit|letter of intent|intent to (bid|propose)",
    "Wet signature / original / notarization": r"notar|\boriginal\b|blue ink|wet signature|corporate seal|sworn",
    "Hard-copy / physical delivery": r"sealed|hard ?cop|envelope|\bmailed\b|hand[- ]deliver",
    "Set-aside / prime eligibility": r"set-aside|small business|prime contractor|\bSPIN\b|E-?Rate",
    "Contract vehicle / approved pool": r"cooperative|contract vehicle|\bGSA\b|approved vendor|master agreement|piggyback",
    "License / certification": r"licen[sc]|certified|certification",
    "Business registration / authority to do business": r"registered|registration|good standing|authori[sz]ed to do business|secretary of state|W-9",
    "Manufacturer / OEM / partner authorization": r"manufacturer|\bOEM\b|authori[sz]ed (reseller|partner|dealer)|or (an )?equivalent",
    "Local presence / geography": r"local (office|vendor|presence)|in-state|within \d+ miles|stocked in",
    "Supplier diversity": r"\b(MBE|WBE|DBE|HUB|SBE|minority|women-owned|supplier diversity)\b",
    "Security / background / hosting / data residency": r"background check|fingerprint|cyber|data residency|hosting|FERPA|security clearance",
    "Conflict / OCI / independence": r"conflict of interest|organizational conflict|independence",
    "Bid bond / bid security": r"bid bond|bid security|certified check|bid guarantee",
    "Performance / payment bond / bonding capacity": r"performance bond|payment bond|\bbond(ed|ing)\b",
    "Financial statements / qualification": r"financial statement|audited|balance sheet|bank (letter|reference)|net worth",
    "Insurance": r"insurance|insured",
    "Budget / ceiling / NTE": r"\bbudget\b|not[- ]to[- ]exceed|\bceiling\b|price limitation",
    "Contract exceptions / redlines / term acceptance": r"exceptions? to|redline|deviation|by submitting|agrees? to (all|the) terms",
    "Contracts objection triggers": r"limitation of liability|indemnif|\bassign|terminat|payment (shall|will) be made|\bnet \d+",
    "Strict staffing / qualifications": r"key personnel|years of experience|\bresumes?\b|staff(ing)? qualification|\bdegree\b",
    "Delivery mix": r"install|configur|\bservices\b|hardware|software|training",
    "Cabling / field installation / site work": r"cabl|\binstall|site work|on-?site",
    "Hardware requirements / quantities": r"quantit|part number|\bequivalent\b",
    "Customer location": r"\b\d+\s+[A-Z][\w.]*\s+(Street|St\.|Road|Rd\.|Avenue|Ave\.)|\b[A-Z]{2}\s+\d{5}\b",
    "Customer profile / scale": r"district-?wide|\bschools\b|\bstudents\b|\bsites\b|\blocations\b|statewide",
}
_SS_FOUND_RISK = {
    "Wet signature / original / notarization": "High",
    "Hard-copy / physical delivery": "High",
    "Set-aside / prime eligibility": "High",
    "Insurance": "Moderate",
    "Contract exceptions / redlines / term acceptance": "High",
    "Contracts objection triggers": "High",
    "Cabling / field installation / site work": "Moderate",
    "Delivery mix": "Informational",
    "Hardware requirements / quantities": "Informational",
    "Customer location": "Informational",
    "Customer profile / scale": "Informational",
}


def _show_stoppers(ctx: Context, tabs) -> list[dict[str, str]]:
    rows = []
    due = next((r for r in tabs["Key Dates"] if r["date_type"] == "Proposal Due Date"), None)
    for area, kind in SHOW_STOPPER_AREAS:
        row = {"screening_area": area, "screening_type": kind, "source_file": ctx.source_file}
        if area == "Quick Turn":
            if due:
                days = int(due["calendar_days_from_shred_run"])
                risk = "Critical" if days <= 7 else "Moderate" if days <= 14 else "None Identified"
                status = "Past Due" if days < 0 else f"{days} calendar days remaining"
                row.update(
                    finding_status="Found",
                    risk_level=risk,
                    finding_requirement=f"Proposal due {due['date']} {due['time']}".strip(),
                    key_value_threshold=status,
                    required_stage="Proposal Submission",
                    source_location=due["source_location"],
                    decision_relevance="Deadline relative to the shred run date.",
                    next_action="Confirm the opportunity is open and the deadline is achievable.",
                )
            else:
                row.update(
                    finding_status="Unclear",
                    risk_level="Watch",
                    finding_requirement="No proposal due date was identified.",
                    next_action="Confirm the proposal deadline.",
                )
            rows.append(row)
            continue
        pattern = _SS_PATTERNS[area]
        hits = ctx.find(pattern, limit=3)
        if hits:
            row.update(
                finding_status="Found",
                risk_level=_SS_FOUND_RISK.get(area, "Watch"),
                decision_relevance=REVIEW,
                finding_requirement=hits[0].text[:400],
                key_value_threshold=f"{len(ctx.find(pattern))} matching statement(s)",
                required_stage=_stage(hits[0].text),
                source_location=hits[0].location,
                related_requirement_ids=ctx.related(pattern),
                next_action="Review the cited language and confirm the risk level.",
            )
        else:
            row.update(
                finding_status="Not Found",
                risk_level="None Identified",
                finding_requirement="No such requirement identified.",
                key_value_threshold="None stated",
                source_location="Entire document reviewed",
                next_action="No action from supplied document; verify any addendum.",
            )
        rows.append(row)
    return rows


def _readiness(ctx: Context, tabs, matrix_text: str | None) -> list[dict[str, str]]:
    has_text = any(t.strip() for t in ctx.result.page_texts)
    rows = [
        {
            "readiness_area": "Overall shred readiness",
            "status": "Partial",
            "source_file_dependency": "Provided procurement package",
            "issue_limitation": "Rule-based extraction: requirements, dates, forms and source "
            "locations are extracted, but risk levels, decision relevance and contradictions are not assessed.",
            "impact_on_findings": "Judgement columns need human review before go/no-go use.",
            "affected_tabs_ids": "Show Stoppers; Contract Objections; Q&A + Issue Log",
            "mitigation_next_step": "Review flagged rows, or re-run with AI analysis enabled.",
        },
        {
            "readiness_area": "RFP source review",
            "status": "Reviewed" if has_text else "Missing Referenced Source",
            "source_file_dependency": ctx.source_file,
            "issue_limitation": "Text layer extracted from all pages."
            if has_text
            else "No text layer (scanned document); OCR is not supported.",
            "affected_tabs_ids": "All tabs",
        },
    ]
    if any(r["file_source_name"].startswith("Addendum") for r in tabs["Source Inventory"]):
        rows.append(
            {
                "readiness_area": "Addendum / Q&A status",
                "status": "Missing Referenced Source",
                "source_file_dependency": "Potential Q&A/addendum release",
                "issue_limitation": "The document references an addendum/Q&A that was not supplied.",
                "impact_on_findings": "An addendum could "
                "alter dates, scope, pricing, eligibility, legal terms, or submission rules.",
                "affected_tabs_ids": "Source Inventory; Amendment + Delta Log; Q&A + Issue Log",
                "mitigation_next_step": "Confirm whether any addendum/Q&A was issued.",
            }
        )
    rows.append(
        {
            "readiness_area": "Contract-objection matrix review",
            "status": "Complete with Limits" if matrix_text else "Missing Referenced Source",
            "source_file_dependency": "Contracts Proposal Objection Matrix",
            "issue_limitation": "Matrix supplied but applied only by AI mode."
            if matrix_text
            else "No objection matrix was uploaded.",
            "affected_tabs_ids": "Contract Objections; Show Stoppers",
            "mitigation_next_step": "Upload the matrix and run AI mode for matrix-based screening.",
        }
    )
    due = next((r for r in tabs["Key Dates"] if r["date_type"] == "Proposal Due Date"), None)
    rows.append(
        {
            "readiness_area": "Deadline status",
            "status": "Reviewed" if due else "Complete with Limits",
            "source_file_dependency": ctx.source_file,
            "issue_limitation": f"Proposal due {due['date']}; {due['calendar_days_from_shred_run']} calendar days "
            "from the shred run date."
            if due
            else "No proposal due date was identified.",
            "affected_tabs_ids": "Key Dates; Show Stoppers SS-001",
        }
    )
    return rows


def _issues(ctx: Context, tabs) -> list[dict[str, str]]:
    rows = []
    due = next((r for r in tabs["Key Dates"] if r["date_type"] == "Proposal Due Date"), None)
    if due and int(due["calendar_days_from_shred_run"]) < 0:
        rows.append(
            {
                "issue_type": "Timing",
                "question_conflict": f"The proposal deadline was {due['date']}, before "
                f"the shred run date of {ctx.run_date:%b %d, %Y}.",
                "severity": "Critical",
                "status": "Open",
                "decision_owner_needed": "Pursuit owner / client procurement contact",
                "impact": "Cannot support an on-time submission without a reopened or reissued opportunity.",
                "source_file": ctx.source_file,
                "source_location": due["source_location"],
                "resolution_notes": "Confirm whether the opportunity was reopened, extended, or reissued.",
            }
        )
    if not due:
        rows.append(
            {
                "issue_type": "Deadline ambiguity",
                "question_conflict": "No proposal due date was identified.",
                "severity": "High",
                "status": "Open",
                "decision_owner_needed": "Procurement contact",
                "source_file": ctx.source_file,
                "resolution_notes": "Confirm the proposal deadline.",
            }
        )
    timed = [r for r in tabs["Key Dates"] if r.get("time")]
    if timed and all(r["time_zone"] == "Not stated" for r in timed):
        rows.append(
            {
                "issue_type": "Deadline ambiguity",
                "question_conflict": "Deadlines state a time but no time zone.",
                "severity": "Watch",
                "status": "Open",
                "decision_owner_needed": "Procurement contact",
                "source_file": ctx.source_file,
                "resolution_notes": "Confirm the controlling time zone.",
            }
        )
    if any(r["readiness_area"] == "Addendum / Q&A status" for r in tabs["Shred Readiness"]):
        rows.append(
            {
                "issue_type": "Missing source",
                "question_conflict": "The document references an addendum/Q&A release that was not supplied.",
                "severity": "High",
                "status": "Open",
                "decision_owner_needed": "Source owner / procurement contact",
                "impact": "A later addendum could change scope, dates, evaluation, pricing or legal terms.",
                "source_file": ctx.source_file,
                "resolution_notes": "Obtain or confirm absence of addenda/Q&A.",
            }
        )
    if not ctx.find(r"limitation of liability", limit=1) and tabs["Contract Objections"]:
        rows.append(
            {
                "issue_type": "Contract acceptance / exceptions",
                "question_conflict": "No limitation-of-liability provision was identified in the terms.",
                "severity": "Watch",
                "status": "Open",
                "decision_owner_needed": "Contracts",
                "source_file": ctx.source_file,
                "resolution_notes": "Complete internal Contracts review.",
            }
        )
    return rows
