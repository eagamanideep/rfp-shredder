"""Shape an ExtractionResult into DataFrames and write them to a formatted Excel workbook."""

from __future__ import annotations

import io
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink

from .extract import ExtractionResult

SHEETS = ("Summary", "Requirements", "Sections", "Tables", "Pages")
COMPLIANCE_OPTIONS = ("Yes", "No", "Partial", "N/A")

_EXCEL_CELL_LIMIT = 32_767
_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_WRAP = Alignment(wrap_text=True, vertical="top")

# Column widths (in characters) for the fixed-layout sheets; anything unlisted gets 18.
_WIDTHS = {
    "Requirement": 80,
    "Content": 100,
    "Text": 110,
    "Section": 45,
    "Path": 50,
    "Heading": 45,
    "Value": 70,
    "Response / Notes": 40,
    "Field": 24,
}


def build_frames(result: ExtractionResult) -> dict[str, pd.DataFrame]:
    requirements = pd.DataFrame(
        [
            {
                "Req ID": r.id,
                "Section": r.section,
                "Page": r.page,
                "Requirement": r.text,
                "Keyword": r.keyword,
                "Type": r.type,
                "Compliance": "",
                "Response / Notes": "",
                "Owner": "",
            }
            for r in result.requirements
        ],
        columns=[
            "Req ID",
            "Section",
            "Page",
            "Requirement",
            "Keyword",
            "Type",
            "Compliance",
            "Response / Notes",
            "Owner",
        ],
    )
    sections = pd.DataFrame(
        [
            {
                "Section ID": s.id,
                "Level": s.level,
                "Heading": s.heading,
                "Path": s.path,
                "Page": s.page,
                "Content": s.content,
                "Words": len(s.content.split()),
            }
            for s in result.sections
        ],
        columns=["Section ID", "Level", "Heading", "Path", "Page", "Content", "Words"],
    )
    pages = pd.DataFrame(
        [{"Page": i, "Text": text, "Characters": len(text)} for i, text in enumerate(result.page_texts, start=1)],
        columns=["Page", "Text", "Characters"],
    )
    tables = pd.DataFrame(
        [
            {
                "Table": table_sheet_name(i, t.page),
                "Page": t.page,
                "Rows": len(t.rows),
                "Columns": len(t.rows[0]),
                "First row": " | ".join(t.rows[0]),
            }
            for i, t in enumerate(result.tables, start=1)
        ],
        columns=["Table", "Page", "Rows", "Columns", "First row"],
    )
    by_type = requirements["Type"].value_counts().to_dict() if not requirements.empty else {}
    summary = pd.DataFrame(
        [
            ("Source file", result.name),
            ("File type", result.file_type.upper()),
            ("Pages", len(result.page_texts) if result.file_type == "pdf" else "n/a"),
            ("Sections", len(result.sections)),
            ("Requirements", len(result.requirements)),
            *((f"  {kind}", count) for kind, count in by_type.items()),
            ("Tables", len(result.tables)),
            ("Generated", datetime.now().strftime("%Y-%m-%d %H:%M")),
            *(("Warning", w) for w in result.warnings),
        ],
        columns=["Field", "Value"],
    )
    return {"Summary": summary, "Requirements": requirements, "Sections": sections, "Tables": tables, "Pages": pages}


def table_sheet_name(index: int, page: int | None) -> str:
    return f"Table {index}" + (f" (p{page})" if page else "")


def build_workbook(result: ExtractionResult, include: tuple[str, ...] = SHEETS) -> bytes:
    frames = build_frames(result)
    wb = Workbook()
    wb.remove(wb.active)

    for name in SHEETS:
        if name not in include:
            continue
        ws = wb.create_sheet(name)
        frame = frames[name]
        _write_rows(ws, [list(frame.columns), *frame.itertuples(index=False, name=None)])
        for idx, column in enumerate(frame.columns, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = _WIDTHS.get(column, 18)
        ws.freeze_panes = "A2"
        if len(frame):
            ws.auto_filter.ref = ws.dimensions

        if name == "Requirements" and len(frame):
            col = get_column_letter(frame.columns.get_loc("Compliance") + 1)
            dv = DataValidation(type="list", formula1='"' + ",".join(COMPLIANCE_OPTIONS) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{col}2:{col}{len(frame) + 1}")
        if name == "Tables":
            for row in range(2, len(frame) + 2):
                cell = ws.cell(row=row, column=1)
                cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{cell.value}'!A1")
                cell.font = Font(color="0563C1", underline="single")

    if "Tables" in include:
        for i, table in enumerate(result.tables, start=1):
            ws = wb.create_sheet(table_sheet_name(i, table.page))
            _write_rows(ws, table.rows)
            for idx in range(1, len(table.rows[0]) + 1):
                ws.column_dimensions[get_column_letter(idx)].width = 30

    if not wb.sheetnames:
        wb.create_sheet("Summary")

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _write_rows(ws, rows: list) -> None:
    for row in rows:
        ws.append([_cell_value(v) for v in row])
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = _WRAP


def _cell_value(value):
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
        if len(value) > _EXCEL_CELL_LIMIT:
            value = value[: _EXCEL_CELL_LIMIT - 20] + " …[truncated]"
    return value
