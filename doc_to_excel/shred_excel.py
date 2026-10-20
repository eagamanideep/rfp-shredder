"""Write a Shred to an .xlsx workbook styled like the Full Shred template."""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .shred import Shred
from .shred_schema import AUDIENCE, TABS, TabSpec

_TITLE_FILL = PatternFill("solid", fgColor="0A0F8C")
_PURPOSE_FILL = PatternFill("solid", fgColor="8DB4E2")
_HEADER_FILL = PatternFill("solid", fgColor="0066CC")
_ID_FILL = PatternFill("solid", fgColor="F2F2F2")
_WHITE_BOLD = Font(bold=True, color="FFFFFF")
_THIN = Side(style="thin", color="000000")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_WRAP_TOP = Alignment(wrap_text=True, vertical="top")
_CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")

# Controlled values -> (fill colour, white bold text?)
_VALUE_COLOURS = {
    "Critical": ("B5479B", True),
    "Partial": ("B5479B", True),
    "High": ("D59BC8", False),
    "Missing Referenced Source": ("D59BC8", False),
    "Moderate": ("E8C6E0", False),
    "Complete with Limits": ("E8C6E0", False),
    "Submission = Acceptance": ("F5E6F1", False),
    "Watch": ("F5E6F1", False),
    "Unclear": ("F5E6F1", False),
    "Informational": ("8DB4E2", False),
    "Reviewed": ("8DB4E2", False),
    "Complete": ("C6E0B4", False),
    "Not Applicable": ("E7E6E6", False),
}
_COLOURED_COLUMNS = {
    "Risk Level",
    "Status",
    "Severity",
    "Confidence",
    "Proposal Risk Level",
    "Client Restriction",
    "Required?",
    "Mandatory?",
    "Requirement Force",
    "Required Stage",
    "Time Zone",
    "Signature Type",
    "Signature / Notary",
}
# Short controlled-value columns are centred, like the template.
_CENTERED_COLUMNS = {
    "Finding Status",
    "Risk Level",
    "Required Stage",
    "Mandatory?",
    "Required?",
    "Quick Turn?",
    "Calendar Days from Shred Run",
    "Time",
    "Time Zone",
    "Quantity",
    "Weight / Points",
    "Procurement Authority?",
    "Reviewed?",
    "Superseded?",
    "Separate Cost File?",
    "Separate Artifact?",
    "File Provided?",
    "Proposal Stage?",
    "Deliverable Stage?",
    "Confidence",
    "Client Restriction",
    "Matrix Escalation",
    "Proposal Risk Level",
    "Copies / Originals",
    "Signature / Notary",
}
_WIDE = {
    "Decision Relevance",
    "Finding / Requirement",
    "Requirement Text",
    "Requirement",
    "Issue / Limitation",
    "Impact on Findings",
    "Question / Conflict",
    "Requirement / Description",
    "Clause Summary",
    "Triggering Language (Short Quote)",
    "Purpose",
    "Key Controlled Values / Notes",
}


def build_shred_workbook(shred: Shred) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for spec in TABS:
        rows = _guide_rows(shred) if spec.name == "Workbook Guide" else shred.tabs.get(spec.name, [])
        _write_tab(wb.create_sheet(spec.name), spec, shred, rows)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _guide_rows(shred: Shred) -> list[dict[str, str]]:
    return [
        {
            "tab": t.name,
            "purpose": t.purpose,
            "primary_use_audience": "Workbook users" if t.name == "Workbook Guide" else AUDIENCE,
            "id_prefix": f"{t.prefix}-",
            "key_controlled_values_notes": shred.guide_notes.get(t.name, t.guide_notes),
        }
        for t in TABS
    ]


def _write_tab(ws, spec: TabSpec, shred: Shred, rows: list[dict[str, str]]) -> None:
    has_id = bool(spec.id_header)
    headers = ([spec.id_header] if has_id else []) + list(spec.headers)
    keys = (["id"] if has_id else []) + list(spec.keys)
    width = len(headers)
    last = get_column_letter(width)

    ws.append([f"{spec.name} | {shred.title_suffix}"])
    ws.append([spec.purpose])
    ws.append([shred.banner_for(spec)])
    ws.append([])
    ws.append(headers)
    for r, fill, font in (
        (1, _TITLE_FILL, Font(bold=True, color="FFFFFF", size=14)),
        (2, _PURPOSE_FILL, Font(bold=True)),
    ):
        ws.merge_cells(f"A{r}:{last}{r}")
        ws[f"A{r}"].fill, ws[f"A{r}"].font = fill, font
    ws.merge_cells(f"A3:{last}3")
    ws["A3"].font = Font(size=9)
    ws["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[3].height = 30
    for cell in ws[5]:
        cell.fill, cell.font, cell.border = _HEADER_FILL, _WHITE_BOLD, _BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[5].height = 30

    for row in rows:
        ws.append([_clean(row.get(k, "")) for k in keys])
        for header, cell in zip(headers, ws[ws.max_row], strict=True):
            cell.border = _BORDER
            cell.alignment = _CENTER if header in _CENTERED_COLUMNS else _WRAP_TOP
            if has_id and cell.column == 1:
                cell.fill, cell.font, cell.alignment = _ID_FILL, Font(bold=True), _CENTER
            colour = _VALUE_COLOURS.get(str(cell.value or "")) if header in _COLOURED_COLUMNS else None
            if colour:
                cell.fill = PatternFill("solid", fgColor=colour[0])
                if colour[1]:
                    cell.font = _WHITE_BOLD

    for idx, header in enumerate(headers, start=1):
        size = 10 if header == spec.id_header and has_id else 40 if header in _WIDE else 20
        ws.column_dimensions[get_column_letter(idx)].width = size
    ws.freeze_panes = "B6" if has_id else "A6"
    if rows:
        ws.auto_filter.ref = f"A5:{last}{ws.max_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:5"


def _clean(value) -> str | int | float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return value
    text = ILLEGAL_CHARACTERS_RE.sub("", str(value)).strip()
    return text[:32_000] or None
