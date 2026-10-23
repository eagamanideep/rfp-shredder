import io
from datetime import date
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from doc_to_excel import extract_document, shred_ai
from doc_to_excel.shred import build_rule_shred
from doc_to_excel.shred_ai import AIConfig, build_ai_shred
from doc_to_excel.shred_excel import build_shred_workbook
from doc_to_excel.shred_schema import SHOW_STOPPER_AREAS, TAB_BY_NAME, TABS, column_key

RFP = """REQUEST FOR PROPOSALS
Network Switch Upgrades
RFP 10-2026-014
SECTION 1 – Instructions
Sealed bids will be received until January 5, 2026 at 2:00 p.m.
Questions may be submitted until December 3, 2025 at 2:00 p.m.
Bids must be signed in blue ink. The Non-Collusion Affidavit must be notarized.
The vendor must provide a W-9.
SECTION 2 – Terms
No right or interest in the contract shall be assigned without written permission.
The contractor shall maintain general liability insurance of $1,000,000.
"""
RUN = date(2026, 9, 20)


def rule_shred(**kwargs):
    result = extract_document(RFP.encode(), "rfp.txt")
    return result, build_rule_shred(result, "rfp.txt", RUN, **kwargs)


def test_rule_shred_fills_every_tab_with_ids():
    _, shred = rule_shred()
    for spec in TABS[1:]:
        rows = shred.tabs[spec.name]
        assert rows, spec.name
        assert rows[0]["id"] == f"{spec.prefix}-001"
    assert shred.rfp_number == "10-2026-014"
    assert [r["screening_area"] for r in shred.tabs["Show Stoppers"]] == [a for a, _ in SHOW_STOPPER_AREAS]


def test_rule_shred_dates_and_quick_turn():
    _, shred = rule_shred()
    dates = {r["date_type"]: r for r in shred.tabs["Key Dates"]}
    assert dates["Shred Run Date"]["id"] == "KD-001"
    due = dates["Proposal Due Date"]
    assert (due["date"], due["calendar_days_from_shred_run"], due["quick_turn"]) == ("Jan 05, 2026", "-258", "Past Due")
    assert dates["Questions Due"]["time_zone"] == "Not stated"
    quick = shred.tabs["Show Stoppers"][0]
    assert (quick["finding_status"], quick["risk_level"]) == ("Found", "Critical")
    assert "Proposal deadline: Past Due" in shred.opportunity_line()


def test_rule_shred_screening_and_forms():
    _, shred = rule_shred()
    stoppers = {r["screening_area"]: r for r in shred.tabs["Show Stoppers"]}
    assert stoppers["Wet signature / original / notarization"]["finding_status"] == "Found"
    assert stoppers["Bid bond / bid security"]["finding_status"] == "Not Found"
    forms = {r["form_attachment"] for r in shred.tabs["Forms + Attachments"]}
    assert {"Non-Collusion Affidavit", "W-9"} <= forms
    assert any(r["trigger_category"] == "Assignment" for r in shred.tabs["Contract Objections"])
    assert shred.tabs["Accessibility 508 ADA"][0]["standard"] == "Not stated"


def test_user_overrides_win():
    _, shred = rule_shred(client="Bartow County School System", title="Switch Upgrades")
    assert shred.title_suffix == "Bartow County School System | RFP 10-2026-014 Switch Upgrades"
    assert shred.client_overridden and not shred.rfp_overridden


def test_workbook_layout_matches_template():
    _, shred = rule_shred(client="Bartow County School System", title="Switch Upgrades")
    wb = load_workbook(io.BytesIO(build_shred_workbook(shred)))
    assert wb.sheetnames == [t.name for t in TABS]
    ws = wb["Show Stoppers"]
    assert ws["A1"].value == "Show Stoppers | Bartow County School System | RFP 10-2026-014 Switch Upgrades"
    assert ws["A2"].value == TAB_BY_NAME["Show Stoppers"].purpose
    assert [c.value for c in ws[5]][:5] == ["ID", "Screening Area", "Screening Type", "Finding Status", "Risk Level"]
    assert ws["A6"].value == "SS-001" and ws["E6"].fill.fgColor.rgb.endswith("B5479B")
    guide = wb["Workbook Guide"]
    assert guide["A3"].value.startswith("Opportunity: Bartow County School System")
    assert [guide.cell(row=r, column=4).value for r in (6, 17)] == ["WG-", "R-"]
    assert wb["Compliance Matrix"]["A5"].value == "Requirement ID"


def test_ai_shred_replaces_tabs_and_falls_back_per_group(monkeypatch):
    result, base = rule_shred()
    rule_issue_log = list(base.tabs["Q&A + Issue Log"])
    rule_cost = list(base.tabs["Cost Breakdown"])
    rule_screening = {n: list(base.tabs[n]) for n in shred_ai.SCREENING_TABS}
    rule_stoppers = [dict(r) for r in base.tabs["Show Stoppers"]]
    issue_notes, stopper_areas, linked = [], [], []

    def fake_call(
        client, config, system, names, context, run_date, opportunity=False, usage=None, excerpt="", note="", areas=None
    ):
        chunk_call = names in shred_ai.CHUNK_CALLS
        if chunk_call:
            assert "EXCERPT" in excerpt and "Sealed bids" in excerpt  # stage 1 reads the document in chunks
            assert "EVALUATION FACTORS" in system and "EVALUATION RUBRIC" in system
        else:
            assert "PROCUREMENT DOCUMENT" in system and "Sealed bids" in system
        if "Show Stoppers" in names:
            stopper_areas.append(list(areas))
            raise RuntimeError("boom")
        if "Q&A + Issue Log" in names:
            issue_notes.append(note)
        if "Cost Breakdown" in names:
            return {column_key("Cost Breakdown"): [], "opportunity": {}}
        if names == shred_ai.SCREENING_TABS:  # nothing to screen; the model reports absence anyway
            not_found = {k: "Not Found in the supplied excerpt" for k in TAB_BY_NAME[names[0]].keys}
            return {column_key(n): [not_found] if n == names[0] else [{}] for n in names} | {"opportunity": {}}
        out = {}
        for name in names:
            row = {k: f"{name} value" for k in TAB_BY_NAME[name].keys}
            out[column_key(name)] = [row, {k: f"{name} second {k}" for k in TAB_BY_NAME[name].keys}]
        if chunk_call:
            assert context == ""
            assert opportunity == (names == ("Key Dates",))  # one stage-1 call reads the opportunity details
            if opportunity:
                out["opportunity"] = {
                    "client": "Bartow County School System",
                    "rfp_number": "X",
                    "title": "Switch Upgrades",
                }
        elif names[0] in shred_ai.STAGE2_TABS:
            # Stage 2 runs alongside stage 1, so it cannot see requirement IDs; links are added afterwards.
            assert context == "" and note == shred_ai.STAGE2_RULE
        else:
            assert "R-001" in context  # stage 3 sees the compliance matrix IDs
        return out

    def fake_link(client, config, shred, tab):
        linked.append(tab)
        return {row["id"]: "R-001" for row in shred.tabs[tab]}

    monkeypatch.setattr(shred_ai, "_call", fake_call)
    monkeypatch.setattr(shred_ai, "_link_call", fake_link)
    progress = []
    shred = build_ai_shred(base, result, AIConfig("sk-test"), progress=lambda f, m: progress.append(f))

    assert shred.mode == "AI (gpt-5.6-sol)"
    assert shred.client == "Bartow County School System" and shred.rfp_number == "X"
    assert [r["id"] for r in shred.tabs["Compliance Matrix"]] == ["R-001", "R-002"]
    assert [r["id"] for r in shred.tabs["Key Dates"]] == ["KD-001", "KD-002", "KD-003"]
    assert shred.tabs["Key Dates"][0]["date_type"] == "Shred Run Date"
    assert shred.tabs["Scope Summary"][0]["scope_area"] == "Scope Summary value"
    # No chunk found anything to screen: the rule-based rows (here Not Found rows) stand, without a warning.
    assert all(shred.tabs[n] == rule_screening[n] for n in shred_ai.SCREENING_TABS)
    assert not any("Screening" in w for w in shred.warnings)
    # An empty answer is a failure too: the rule-based rows stay and the run says so.
    assert shred.tabs["Cost Breakdown"] == rule_cost and any("no rows for Cost Breakdown" in w for w in shred.warnings)
    # Show Stoppers runs in parts covering every fixed area once; when all parts fail, each area keeps its rule row.
    assert sorted(i for part in stopper_areas for i in part) == list(range(len(SHOW_STOPPER_AREAS)))
    assert any("Show Stoppers" in w for w in shred.warnings)
    assert [r["finding_status"] for r in shred.tabs["Show Stoppers"]] == [r["finding_status"] for r in rule_stoppers]
    # The issue log runs once per issue kind and its parts are merged without repeats.
    assert sorted(issue_notes) == sorted(shred_ai.ISSUE_LOG_PARTS)
    assert len(shred.tabs["Q&A + Issue Log"]) == 2
    # Link calls fill the requirement-ID columns of the tabs built before the matrix.
    assert (
        "Submission Mechanics" in linked and shred.tabs["Submission Mechanics"][0]["related_requirement_ids"] == "R-001"
    )
    # Constant columns are filled in code.
    assert shred.tabs["Compliance Matrix"][0]["source_file"] == "Compliance Matrix value"  # the model's value is kept
    # The other stage-3 calls still land.
    assert shred.tabs["Q&A + Issue Log"] != rule_issue_log
    assert shred.tabs["Q&A + Issue Log"][0]["severity"] == "Q&A + Issue Log value"
    assert len(shred.tabs["Show Stoppers"]) == len(SHOW_STOPPER_AREAS)
    assert progress[-1] == 1.0

    # Every step is timed, failed groups included.
    steps = [step for step, _ in shred.timings]
    assert steps[0] == "Rule-based shred (all tabs)"
    assert "AI stage 1 · Compliance Matrix · p. 1" in steps and "AI stage 1 · Key Dates · p. 1" in steps
    assert "AI stage 1 total (chunks run in parallel)" in steps
    assert "AI stage 2 · Scope Summary" in steps and "AI stage 2 · Amendment + Delta Log" in steps
    assert "AI stage 2 total (groups run in parallel)" in steps
    assert "AI stages 1+2 total (run together)" in steps
    assert "AI stage 3 · Show Stoppers 1/3" in steps and "AI stage 3 · Q&A + Issue Log 2/2" in steps
    assert "AI stage 3 · Links · Submission Mechanics" in steps
    assert "AI stage 3 total (groups run in parallel)" in steps
    assert all(seconds >= 0 for _, seconds in shred.timings)


def test_run_log_records_stage_times(tmp_path):
    from doc_to_excel.pipeline import generate_full_shred
    from doc_to_excel.run_log import load_runs, record_run

    result = extract_document(RFP.encode(), "rfp.txt")
    shred, workbook, total = generate_full_shred(result, ("Read & extract document (TXT)", 1.5), "rfp.txt", RUN)
    assert workbook[:2] == b"PK"  # a real .xlsx zip
    assert total >= 1.5
    assert [s for s, _ in shred.timings] == [
        "Read & extract document (TXT)",
        "Rule-based shred (all tabs)",
        "Build Excel workbook",
    ]

    log = tmp_path / "times.csv"
    record_run(shred, total, 1, path=log)
    shred.timings.append(("AI stage 2 · Scope Summary", 9.0))
    shred.timings.append(("AI stage 2 total (groups run in parallel)", 12.0))
    record_run(shred, total + 20, 1, model="gpt-6.1-sol", reasoning_effort="high", path=log)

    rule, ai = load_runs(log)
    assert (rule["mode"], rule["model"], rule["read_extract_s"], rule["ai_stage1_s"]) == ("Rule-based", "", "1.50", "")
    assert (ai["mode"], ai["model"], ai["reasoning_effort"]) == ("AI", "gpt-6.1-sol", "high")
    assert ai["ai_stage2_s"] == "12.00"  # the parallel stage's wall-clock total, not one group


def test_cost_uses_cached_and_long_context_rates():
    # gpt-5.6-sol: $4 input, $0.40 cached, $20 output per 1M; over 272K input: $8 / $0.80 / $30.
    assert shred_ai.cost_usd("gpt-5.6-sol", 100_000, 40_000, 10_000) == pytest.approx(0.24 + 0.016 + 0.2)
    assert shred_ai.cost_usd("gpt-5.6-sol", 300_000, 0, 10_000) == pytest.approx(2.4 + 0.3)
    assert shred_ai.cost_usd("gpt-5", 300_000, 0, 0) == pytest.approx(0.375)  # no long-context tier
    assert shred_ai.cost_usd("unknown-model", 1, 0, 1) is None
    assert set(shred_ai.PRICES) == set(shred_ai.MODELS)


def test_call_records_usage_even_when_response_is_cut_off():
    usage = SimpleNamespace(
        prompt_tokens=1_000,
        completion_tokens=500,
        prompt_tokens_details=SimpleNamespace(cached_tokens=200),
        completion_tokens_details=SimpleNamespace(reasoning_tokens=300),
    )
    message = SimpleNamespace(content="{}", refusal=None)
    response = SimpleNamespace(usage=usage, choices=[SimpleNamespace(finish_reason="length", message=message)])
    client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=lambda **_: response)))

    recorded = []
    with pytest.raises(shred_ai.ShredAIError):
        shred_ai._call(client, AIConfig("sk-test"), "sys", ("Key Dates",), "", "2026-09-20", usage=recorded)
    assert recorded == [
        {
            "step": "Key Dates",
            "input_tokens": 1_000,
            "cached_input_tokens": 200,
            "output_tokens": 500,
            "reasoning_tokens": 300,
            "cost_usd": pytest.approx((800 * 4 + 200 * 0.4 + 500 * 20) / 1e6),
        }
    ]


def test_run_log_records_usage_cost_and_citations(tmp_path):
    from doc_to_excel.run_log import citation_stats, load_runs, record_run

    _, shred = rule_shred()
    for name in shred.tabs:
        shred.tabs[name] = []
    shred.tabs["Compliance Matrix"] = [
        {"source_location": "p. 2, Section 1"},
        {"source_location": "p. 9, Section 4"},  # page past the end of a 3-page document
        {"source_location": "Entire document reviewed"},
        {"source_location": ""},
    ]
    assert citation_stats(shred, 3) == {"rows": 4, "cited_rows": 3, "cited_pct": 75.0, "page_cited_pct": 25.0}
    assert citation_stats(shred, None)["page_cited_pct"] is None

    shred.usage = [
        {
            "step": "a",
            "input_tokens": 10,
            "cached_input_tokens": 4,
            "output_tokens": 5,
            "reasoning_tokens": 2,
            "cost_usd": 0.5,
        },
        {
            "step": "b",
            "input_tokens": 20,
            "cached_input_tokens": 0,
            "output_tokens": 7,
            "reasoning_tokens": 1,
            "cost_usd": 0.25,
        },
    ]
    log = tmp_path / "times.csv"
    log.write_text("timestamp,file\n2026-01-01T00:00:00,old.pdf\n")  # a log from before the new columns
    record_run(shred, 12.0, 3, model="gpt-5.6-sol", path=log)

    old, new = load_runs(log)
    assert old["file"] == "old.pdf" and old["cost_usd"] == ""
    assert (new["input_tokens"], new["cached_input_tokens"], new["output_tokens"]) == ("30", "4", "12")
    assert (new["cost_usd"], new["cited_rows"], new["cited_pct"], new["page_cited_pct"]) == (
        "0.7500",
        "3/4",
        "75.0",
        "25.0",
    )


def test_every_call_sends_the_same_schema():
    """The prompt cache covers the response schema, so it must not change between calls."""
    sent = []

    def create(**kwargs):
        sent.append(kwargs)
        message = SimpleNamespace(content="{}", refusal=None)
        return SimpleNamespace(usage=None, choices=[SimpleNamespace(finish_reason="stop", message=message)])

    client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=create)))
    for names in (("Compliance Matrix",), ("Key Dates",), ("Show Stoppers",)):
        shred_ai._call(client, AIConfig("sk-test"), "sys", names, "", "2026-09-20", names == ("Key Dates",))
    formats = [kwargs["response_format"] for kwargs in sent]
    assert formats[0] == formats[1] == formats[2]
    schema = formats[0]["json_schema"]["schema"]
    assert set(schema["required"]) == {column_key(n) for n in shred_ai.AI_TABS} | {"opportunity"}
    assert "Produce ONLY these tabs" in sent[2]["messages"][1]["content"]


def _pdf_like(pages, sections=()):
    from doc_to_excel.extract import ExtractionResult, Section

    heads = [Section(i, 1, h, h, page) for i, (h, page) in enumerate(sections)]
    return ExtractionResult("rfp.pdf", "pdf", list(pages), heads, [], [])


def test_chunks_cut_before_sections_and_show_the_next_page_as_context():
    from doc_to_excel.chunking import plan_chunks

    result = _pdf_like([f"page {i} " + "x" * 1000 for i in range(1, 11)], sections=[("SECTION 2", 5)])
    chunks = plan_chunks(result, 4500)
    assert [(c.first_page, c.last_page) for c in chunks] == [(1, 4), (5, 8), (9, 10)]
    assert "=== Page 4 ===" in chunks[0].text and "=== Page 5 ===" not in chunks[0].text
    assert "Start of page 5: CONTEXT ONLY" in chunks[0].text and "page 5 x" in chunks[0].text
    assert "CONTEXT ONLY" not in chunks[-1].text
    assert chunks[1].label == "pp. 5–8"
    assert [(c.first_page, c.last_page) for c in plan_chunks(result, 0)] == [(1, 10)]


def test_merge_drops_repeats_and_rows_from_the_lookahead():
    from doc_to_excel.chunking import Chunk, merge_chunk_rows

    text = lambda r: r["requirement_text"]  # noqa: E731
    a = {"requirement_text": "The vendor shall submit three signed copies", "source_location": "p. 1"}
    b = {"requirement_text": "Proposals must be delivered by the due date", "source_location": "p. 2, Section 1"}
    lookahead = {"requirement_text": "The contractor shall maintain insurance", "source_location": "p. 3"}
    marker = {"requirement_text": "Bids are opened publicly", "source_location": "p. 2, CONTEXT ONLY"}
    repeat_of_b = {"requirement_text": "Proposals must be delivered by the due date.", "source_location": "p. 4"}
    c = {"requirement_text": "The contractor shall maintain insurance", "source_location": "p. 3"}
    p1 = {"requirement_text": "Priority 1 is worth 25 points", "source_location": "p. 4"}
    p2 = {"requirement_text": "Priority 2 is worth 25 points", "source_location": "p. 4"}
    chunks = [Chunk(1, 2, ""), Chunk(3, 4, "")]
    merged = merge_chunk_rows([[a, b, lookahead, marker], [repeat_of_b, c, p1, p2]], text, chunks)
    assert merged == [a, b, c, p1, p2]
    assert b["source_location"] == "p. 2, Section 1; also p. 4"


def test_citation_check_fixes_the_page_or_flags_the_row():
    from doc_to_excel.chunking import check_citations

    result = _pdf_like(
        ["Cover page for the solicitation", "The vendor shall submit three signed copies of the proposal by mail"]
    )
    right = {"requirement_text": "The vendor shall submit three signed copies", "source_location": "p. 2, Section 1"}
    wrong = {"requirement_text": "The vendor shall submit three signed copies", "source_location": "p. 1, Section 1"}
    made_up = {"requirement_text": "The vendor shall provide a performance bond", "source_location": "p. 2"}
    counts = check_citations([right, wrong, made_up], result, "requirement_text")
    assert counts == {"checked": 3, "corrected": 1, "unverified": 1}
    assert right["source_location"] == "p. 2, Section 1" and "notes" not in right
    assert wrong["source_location"] == "p. 2, Section 1" and "corrected to p. 2" in wrong["notes"]
    assert made_up["notes"].startswith("Unverified")


def test_key_dates_are_normalised_and_day_counts_computed():
    row = {"date": "September 4, 2026", "calendar_days_from_shred_run": "-50"}
    shred_ai._normalize_date(row, date(2026, 10, 22))
    assert row == {"date": "Sep 4, 2026", "calendar_days_from_shred_run": "-48"}
    unparsed = {"date": "Within 15 days of contract approval", "calendar_days_from_shred_run": ""}
    shred_ai._normalize_date(unparsed, date(2026, 10, 22))
    assert unparsed["date"] == "Within 15 days of contract approval"


def test_related_ids_link_requirements_to_stated_dates_and_named_forms():
    _, shred = rule_shred()
    shred.tabs["Key Dates"] = [
        {"id": "KD-001", "date": "Sep 20, 2026"},
        {"id": "KD-002", "date": "Jan 5, 2026"},
        {"id": "KD-003", "date": "Dec 3, 2025"},
    ]
    shred.tabs["Forms + Attachments"] = [
        {"id": "FA-001", "form_attachment": "Non-Collusion Affidavit (notarized)"},
        {"id": "FA-002", "form_attachment": "W-9"},  # too short a name to match safely
    ]
    shred.tabs["Compliance Matrix"] = [
        {"requirement_text": "Sealed bids will be received until January 5, 2026 at 2:00 p.m.", "related_ids": ""},
        {"requirement_text": "The Non-Collusion Affidavit must be notarized.", "related_ids": ""},
        {"requirement_text": "The vendor must provide a W-9.", "related_ids": ""},
        {"requirement_text": "Questions until December 3, 2025.", "related_ids": "R-009"},  # model's link kept
    ]
    shred_ai._link_related_ids(shred)
    assert [r["related_ids"] for r in shred.tabs["Compliance Matrix"]] == ["KD-002", "FA-001", "", "R-009"]


def test_code_filled_columns_are_left_out_of_the_schema_and_filled_after():
    schema = shared_schema_items("Compliance Matrix")
    assert "source_file" not in schema and "amendment_version" not in schema and "requirement_text" in schema
    _, shred = rule_shred(rfp_number="10-2026-014")
    shred.tabs["Compliance Matrix"] = [{"requirement_text": "x", "source_location": "p. 1"}]
    shred.tabs["Scope Summary"] = [{"scope_area": "y"}]
    shred_ai._fill_code_columns(shred)
    assert shred.tabs["Compliance Matrix"][0]["source_file"] == "rfp.txt"
    assert shred.tabs["Compliance Matrix"][0]["amendment_version"] == "Baseline 10-2026-014"
    assert shred.tabs["Scope Summary"][0]["source_file"] == "rfp.txt"


def shared_schema_items(name):
    return shred_ai.shared_schema()["properties"][column_key(name)]["items"]["properties"]


def test_link_call_keeps_only_requirement_ids_that_exist():
    import json

    _, shred = rule_shred()
    shred.tabs["Compliance Matrix"] = [{"id": "R-001", "requirement_text": "Sign in blue ink"}]
    shred.tabs["Submission Mechanics"] = [{"id": "SM-001", "control": "Signature"}, {"id": "SM-002", "control": "x"}]
    reply = {"links": [{"id": "SM-001", "requirement_ids": "R-001, R-999"}]}  # R-999 does not exist; SM-002 omitted
    message = SimpleNamespace(content=json.dumps(reply), refusal=None)
    response = SimpleNamespace(usage=None, choices=[SimpleNamespace(finish_reason="stop", message=message)])
    sent = []
    create = lambda **kwargs: sent.append(kwargs) or response  # noqa: E731
    client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=create)))

    links = shred_ai._link_call(client, AIConfig("sk-test"), shred, "Submission Mechanics")
    assert links == {"SM-001": "R-001", "SM-002": ""}
    assert "related_requirement_ids" not in shred.tabs["Submission Mechanics"][0]  # applied by the caller
    shred_ai._apply_links(shred, "Submission Mechanics", links)
    assert [r["related_requirement_ids"] for r in shred.tabs["Submission Mechanics"]] == ["R-001", ""]
    assert "R-001: Sign in blue ink" in sent[0]["messages"][0]["content"]


def test_link_call_retries_when_no_row_gets_a_link():
    import json

    _, shred = rule_shred()
    shred.tabs["Compliance Matrix"] = [{"id": "R-001", "requirement_text": "Submit the W-9"}]
    shred.tabs["Forms + Attachments"] = [{"id": f"FA-00{i}", "form_attachment": "W-9"} for i in range(1, 4)]
    replies = [{"links": []}, {"links": [{"id": "FA-001", "requirement_ids": "R-001"}]}]

    def create(**_):
        message = SimpleNamespace(content=json.dumps(replies.pop(0)), refusal=None)
        return SimpleNamespace(usage=None, choices=[SimpleNamespace(finish_reason="stop", message=message)])

    client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=create)))
    links = shred_ai._link_call(client, AIConfig("sk-test"), shred, "Forms + Attachments")
    assert list(links.values()) == ["R-001", "", ""]
    assert replies == []  # the empty first answer was retried


def test_run_parallel_resends_a_straggler_and_keeps_the_first_copy(monkeypatch):
    import threading

    monkeypatch.setattr(shred_ai, "STRAGGLER_MIN_SECONDS", 0.2)
    calls = {"slow": 0}
    release = threading.Event()

    def slow():
        calls["slow"] += 1
        if calls["slow"] == 1:
            release.wait(5)  # the first copy stalls
            return "late"
        return "resent"

    jobs = {"a": lambda: "a", "b": lambda: "b", "c": lambda: "c", "slow": slow}
    results = shred_ai._run_parallel(jobs, lambda _key: "same kind")
    release.set()
    assert {k: v[0] for k, v in results.items()} == {"a": "a", "b": "b", "c": "c", "slow": "resent"}
    assert results["slow"][3] and not results["a"][3]  # the resent flag
    assert calls["slow"] == 2


def test_run_parallel_does_not_resend_without_peers_or_after_a_failure(monkeypatch):
    import time

    monkeypatch.setattr(shred_ai, "STRAGGLER_MIN_SECONDS", 0.1)
    sent = []

    def lone():
        sent.append("lone")
        time.sleep(0.5)
        return "lone"

    def fails():
        raise RuntimeError("boom")

    results = shred_ai._run_parallel({"lone": lone, "x": fails}, lambda key: key)  # each in its own kind
    assert results["lone"][:2] == ("lone", None) and sent == ["lone"]
    assert isinstance(results["x"][1], RuntimeError)


def _docx_like(sections, tables=()):
    from doc_to_excel.extract import ExtractionResult, Paragraph, Section, Table

    heads = [
        Section(i, level, heading, heading, None, [Paragraph(text, None) for text in paragraphs])
        for i, (level, heading, paragraphs) in enumerate(sections, start=1)
    ]
    text = "\n".join(f"{h.heading}\n{h.content}" for h in heads)
    return ExtractionResult("rfp.docx", "docx", [text], heads, [], [Table(None, rows) for rows in tables])


def test_word_documents_are_chunked_by_section():
    from doc_to_excel.chunking import plan_chunks

    body = "The vendor shall comply. " * 80  # ~2,000 characters per paragraph
    result = _docx_like(
        [
            (1, "1. INTRODUCTION", [body]),
            (2, "1.1 Background", [body]),
            (1, "2. SCOPE OF WORK", [body, body, body]),  # longer than a chunk: split at paragraphs
            (1, "3. SUBMISSION", [body]),
        ],
        tables=[[["Item", "Qty"], ["Switch", "40"]]],
    )
    chunks = plan_chunks(result, 4500)
    assert all(c.first_page == c.last_page == 0 for c in chunks)
    assert chunks[0].headings == ("1. INTRODUCTION", "1.1 Background")
    assert "2. SCOPE OF WORK" in chunks[1].text and "CONTEXT ONLY" in chunks[0].text
    assert "no page numbers" in chunks[0].text
    assert sum(c.text.count(body) for c in chunks[:-1]) >= 6  # every paragraph belongs to a chunk
    assert chunks[-1].headings == ("Tables",) and "Switch | 40" in chunks[-1].text
    assert chunks[1].label.startswith("§ 2. SCOPE OF WORK")
    # A chunk owns the rule-based rows that cite its headings (used when its call fails).
    assert chunks[0].owns("1.1 Background") and not chunks[0].owns("3. SUBMISSION")
    # A short document, or chunking turned off, stays one chunk.
    assert len(plan_chunks(_docx_like([(1, "A", ["short"])]), 4500)) == 1
    assert len(plan_chunks(result, 0)) == 1 and "Switch | 40" in plan_chunks(result, 0)[0].text


def test_merge_keeps_section_chunk_rows_that_cite_pages():
    from doc_to_excel.chunking import Chunk, merge_chunk_rows

    rows = [[{"requirement_text": "Submit two copies", "source_location": "p. 4, 3. SUBMISSION"}]]
    merged = merge_chunk_rows(rows, lambda r: r["requirement_text"], [Chunk(0, 0, "", ("3. SUBMISSION",))])
    assert len(merged) == 1  # no pages to compare against, so the row is kept


def test_stage3_context_goes_in_the_system_message_so_the_cache_covers_it():
    sent = []

    def create(**kwargs):
        sent.append(kwargs)
        message = SimpleNamespace(content="{}", refusal=None)
        return SimpleNamespace(usage=None, choices=[SimpleNamespace(finish_reason="stop", message=message)])

    client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=create)))
    for names in (("Shred Readiness",), ("Q&A + Issue Log",)):
        shred_ai._call(client, AIConfig("sk-test"), "sys", names, "R-001: Sign", "2026-09-20")
    systems = [kwargs["messages"][0]["content"] for kwargs in sent]
    users = [kwargs["messages"][1]["content"] for kwargs in sent]
    assert systems[0] == systems[1] and systems[0].startswith("sys") and "R-001: Sign" in systems[0]
    assert users[0] != users[1] and not any("R-001" in u for u in users)


def test_evaluation_rubric_shows_the_pages_the_factors_cite():
    _, shred = rule_shred()
    result = _pdf_like(["Cover page", "Technical approach 60 points. Price 40 points.", "Terms", "Scoring steps"])
    shred.tabs["Evaluation Summary"] = [
        {"criterion_factor": "Technical approach", "source_location": "p. 2, table"},
        {"criterion_factor": "Price", "source_location": "p. 4, Evaluation"},
    ]
    rubric = shred_ai._evaluation_rubric(shred, result)
    assert "=== Page 2 ===\nTechnical approach 60 points" in rubric and "=== Page 4 ===" in rubric
    assert "Cover page" not in rubric and "=== Page 3 ===\nTerms" in rubric  # between two cited pages
    assert shred_ai._evaluation_rubric(shred, result, limit=20).endswith("omitted]")
    # Pages between cited evaluation pages are filled in when close together (the rubric sits between them).
    pages = _pdf_like([f"page {i}" for i in range(1, 21)])
    shred.tabs["Evaluation Summary"] = [{"source_location": f"p. {p}"} for p in (3, 6, 15)]
    assert [p for p in range(1, 21) if f"=== Page {p} ===" in shred_ai._evaluation_rubric(shred, pages)] == [
        3,
        4,
        5,
        6,
        15,
    ]
    shred.tabs["Evaluation Summary"] = []
    assert shred_ai._evaluation_rubric(shred, result) == "(not found)"


def test_evaluation_rubric_without_pages_uses_the_cited_sections():
    _, shred = rule_shred()
    result = _docx_like([(1, "4. EVALUATION", ["Technical 60 points"]), (1, "5. TERMS", ["Net 30"])])
    shred.tabs["Evaluation Summary"] = [{"criterion_factor": "Technical", "source_location": "4. EVALUATION"}]
    assert shred_ai._evaluation_rubric(shred, result) == "4. EVALUATION\nTechnical 60 points"


def test_second_look_places_drops_or_keeps_unverified_rows():
    from doc_to_excel.chunking import UNVERIFIED

    result = _pdf_like(["Cover", "Bids are due", "Vendors furnish proof of coverage", "Terms", "Signatures"])
    rows = [
        {"requirement_text": "Verified elsewhere", "source_location": "p. 1", "notes": ""},
        {"requirement_text": "Proof of insurance is needed", "source_location": "p. 2, Section 1", "notes": UNVERIFIED},
        {
            "requirement_text": "A performance bond is needed",
            "source_location": "p. 4",
            "notes": f"Model note; {UNVERIFIED}",
        },
        {"requirement_text": "Call fails", "source_location": "p. 5", "notes": UNVERIFIED},
        {"requirement_text": "Model names a page it was not shown", "source_location": "p. 1", "notes": UNVERIFIED},
    ]
    answers = {"Proof": 3, "A performance": 0, "Model": 5}
    shown = {}

    def recheck(row, pages):
        shown[row["requirement_text"]] = pages
        if row["requirement_text"] == "Call fails":
            raise RuntimeError("boom")
        return next(page for start, page in answers.items() if row["requirement_text"].startswith(start))

    kept, placed, dropped = shred_ai._second_look(rows, result, recheck, 8)
    assert "Verified elsewhere" not in shown  # only Unverified rows are checked
    assert {1, 2, 3} <= set(shown["Proof of insurance is needed"])  # the cited page and its neighbours
    assert placed == 1 and [r["requirement_text"] for r in dropped] == ["A performance bond is needed"]
    assert len(kept) == 4 and rows[2] not in kept
    assert rows[1]["source_location"] == "p. 3, Section 1"
    assert rows[1]["notes"].startswith("Placed on p. 3") and UNVERIFIED not in rows[1]["notes"]
    assert rows[3]["notes"] == rows[4]["notes"] == UNVERIFIED  # a failed call or an unshown page changes nothing


def test_recheck_call_uses_the_small_model_and_shows_the_pages():
    import json

    sent = []

    def create(**kwargs):
        sent.append(kwargs)
        message = SimpleNamespace(content=json.dumps({"page": 3}), refusal=None)
        return SimpleNamespace(usage=None, choices=[SimpleNamespace(finish_reason="stop", message=message)])

    client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=create)))
    row = {"requirement_text": "Proof of insurance", "source_location": "p. 2"}
    config = AIConfig("sk-test", model="gpt-6-luna", reasoning_effort="medium")
    assert shred_ai._recheck_call(client, config, row, [2, 3], ["one", "two", "three"], []) == 3
    assert (sent[0]["model"], sent[0]["reasoning_effort"]) == ("gpt-5-nano", "low")
    user = sent[0]["messages"][1]["content"]
    assert "=== Page 2 ===\ntwo" in user and "=== Page 3 ===\nthree" in user and "one" not in user
    assert config.model == "gpt-6-luna"  # the shred's own config is untouched


def test_force_check_notes_disagreements_without_changing_the_label():
    rows = [
        {"requirement_force": "Mandatory", "requirement_text": "The vendor shall submit three copies."},
        {"requirement_force": "Mandatory", "requirement_text": "Vendors may attend the site visit."},
        {"requirement_force": "Advisory", "requirement_text": "Bids must be signed in blue ink.", "notes": "x"},
        {"requirement_force": "Mandatory", "requirement_text": "If bonding is required, submit a bid bond."},
        {"requirement_force": "Mandatory", "requirement_text": "Bids are due May 5, 2026 at 2:00 PM."},
        {"requirement_force": "Mandatory", "requirement_text": "Report barriers that may affect delivery."},
        {"requirement_force": "Mandatory", "requirement_text": "Payments may not exceed the Price Limitation."},
        {"requirement_force": "Informational", "requirement_text": "The Department may reject any bid."},
        {
            "requirement_force": "Conditional",
            "requirement_text": "Submit references if requested; they must be recent.",
        },
    ]
    assert shred_ai._check_force(rows) == 3
    assert "notes" not in rows[0] and "notes" not in rows[4] and "notes" not in rows[5] and "notes" not in rows[6]
    assert rows[1]["notes"] == 'Force check: the wording ("may") reads as Advisory, not Mandatory.'
    assert rows[2]["notes"] == 'x; Force check: the wording ("must") reads as Mandatory, not Advisory.'
    assert "reads as Conditional" in rows[3]["notes"]
    assert [r["requirement_force"] for r in rows][:4] == ["Mandatory", "Mandatory", "Advisory", "Mandatory"]
