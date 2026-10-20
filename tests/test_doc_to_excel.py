import io

import docx
from openpyxl import load_workbook

from doc_to_excel import build_workbook, extract_document
from doc_to_excel.extract import _merge_split_columns, clean_table_rows, split_sentences

SAMPLE = """SECTION 1 – Overview
A. Scope
The vendor must deliver a work plan within 15 days. Reports should be monthly.
The vendor shall:
o Attend coordination meetings;
o Keep attendance records.
Should a vendor fail to submit, the proposal is rejected.

SECTION 2 – Budget
1. Payment Structure
Payments will be made per deliverable.
"""


def extract_text(text: str, **kwargs):
    return extract_document(text.encode(), "sample.txt", **kwargs)


def test_sections_follow_heading_hierarchy():
    result = extract_text(SAMPLE)
    paths = [s.path for s in result.sections]
    assert paths == [
        "SECTION 1 – Overview",
        "SECTION 1 – Overview > A. Scope",
        "SECTION 2 – Budget",
        "SECTION 2 – Budget > 1. Payment Structure",
    ]


def test_requirements_classified_and_list_items_inherited():
    result = extract_text(SAMPLE)
    reqs = {r.text: (r.keyword, r.type) for r in result.requirements}
    assert reqs["The vendor must deliver a work plan within 15 days."] == ("must", "Mandatory")
    assert reqs["Reports should be monthly."] == ("should", "Recommended")
    assert reqs["o Attend coordination meetings;"] == ("shall (list item)", "Mandatory")
    assert reqs["o Keep attendance records."] == ("shall (list item)", "Mandatory")
    # Conditional "Should a vendor ..." is not a recommendation, and "will" is not a default keyword.
    assert not any(r.text.startswith("Should a vendor") for r in result.requirements)
    assert not any("Payments will" in r.text for r in result.requirements)


def test_list_item_inheritance_can_be_disabled():
    result = extract_text(SAMPLE, inherit_list_items=False)
    assert not any(r.keyword.endswith("(list item)") for r in result.requirements)


def test_sentence_split_keeps_abbreviations():
    text = "Governed by N.H. RSA chapter 91-A. Vendors must comply, e.g. with insurance rules."
    assert split_sentences(text) == [
        "Governed by N.H. RSA chapter 91-A.",
        "Vendors must comply, e.g. with insurance rules.",
    ]


def test_table_cleanup_merges_split_cells():
    raw = [
        ["", "EVENT", None, "", "DATE", None],
        ["RFP Released", None, None, "July 17", None, None],
        ["Proposal Due", None, None, "Sept 4", None, None],
        ["Selection", None, None, "Sept 25", None, None],
    ]
    assert clean_table_rows(raw) == [
        ["EVENT", "DATE"],
        ["RFP Released", "July 17"],
        ["Proposal Due", "Sept 4"],
        ["Selection", "Sept 25"],
    ]


def test_sparse_tick_columns_are_not_merged():
    rows = [["Item", "Yes", "No"], ["A", "x", ""], ["B", "", "x"], ["C", "x", ""], ["D", "", "x"]]
    assert _merge_split_columns(rows) == rows


def test_docx_headings_and_tables():
    document = docx.Document()
    document.add_heading("Security Requirements", level=1)
    document.add_paragraph("The system must support SAML 2.0 single sign-on.")
    document.add_heading("Session Management", level=2)
    document.add_paragraph("Sessions shall expire after 30 minutes of inactivity.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "ID", "Requirement"
    table.cell(1, 0).text, table.cell(1, 1).text = "SSO-1", "Support MFA"
    buffer = io.BytesIO()
    document.save(buffer)

    result = extract_document(buffer.getvalue(), "spec.docx")
    assert [s.path for s in result.sections] == [
        "Security Requirements",
        "Security Requirements > Session Management",
    ]
    assert [r.section for r in result.requirements] == [s.path for s in result.sections]
    assert result.tables[0].rows == [["ID", "Requirement"], ["SSO-1", "Support MFA"]]


def test_workbook_has_expected_sheets_and_compliance_dropdown():
    document = docx.Document()
    document.add_paragraph("The vendor must comply.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "a", "b"
    buffer = io.BytesIO()
    document.save(buffer)
    result = extract_document(buffer.getvalue(), "x.docx")

    wb = load_workbook(io.BytesIO(build_workbook(result)))
    assert wb.sheetnames == ["Summary", "Requirements", "Sections", "Tables", "Pages", "Table 1"]
    ws = wb["Requirements"]
    assert ws["D2"].value == "The vendor must comply."
    assert ws.data_validations.dataValidation[0].formula1 == '"Yes,No,Partial,N/A"'
    assert wb["Tables"]["A2"].hyperlink.location == "'Table 1'!A1"


def test_workbook_respects_sheet_selection():
    result = extract_text(SAMPLE)
    wb = load_workbook(io.BytesIO(build_workbook(result, ("Requirements",))))
    assert wb.sheetnames == ["Requirements"]
