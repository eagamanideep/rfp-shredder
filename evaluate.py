"""Score Full Shred workbooks against the source PDF and a reference workbook. Makes no model calls.

Usage:
    python evaluate.py RFP.pdf REFERENCE.xlsx CANDIDATE.xlsx [CANDIDATE.xlsx ...]

Compliance Matrix measures, per workbook:
- grounded: rows whose content words are at least 80% present on the cited page(s), page text plus table cells. This
  is the pipeline's own citation check (chunking.check_citations), so a row the second look placed on a page but
  worded differently from it still counts as ungrounded.
- requirement recall: reference requirements matched by some row (at least 60% of the reference row's content words
  in one row).
- date recall: reference Key Dates whose date appears in the workbook's Key Dates.
- eval link / related: rows with a non-empty Evaluation Link / Related IDs.

The reference so far is model output (a gpt-5.6-sol workbook), so recall measures agreement with it, not with the
document; a hand-checked answer key would replace it (backlog item 10).
"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

from doc_to_excel import extract_document
from doc_to_excel.chunking import FOUND_ON_PAGE, cited_pages, content_words, page_texts_with_tables
from doc_to_excel.extract import ExtractionResult
from doc_to_excel.shred_ai import DATE_IN_TEXT, _parse_date

# Share of a reference requirement's content words that one row must contain to match it.
MATCHES_REFERENCE = 0.6


def load_tab(path: Path, name: str) -> list[dict]:
    """A workbook tab's rows as {header: value}, from the header row (the one whose first cell ends in "ID")."""
    rows = list(load_workbook(path, read_only=True)[name].iter_rows(values_only=True))
    head = next(i for i, row in enumerate(rows) if row and row[0] and str(row[0]).endswith("ID"))
    return [dict(zip(rows[head], row, strict=False)) for row in rows[head + 1 :] if row and row[0]]


def dates(rows: list[dict]) -> set:
    return {d for r in rows for d in map(_parse_date, DATE_IN_TEXT.findall(str(r.get("Date") or ""))) if d}


def score(result: ExtractionResult, reference: Path, path: Path) -> dict[str, float | int]:
    page_words = [set(content_words(p)) for p in page_texts_with_tables(result)]
    rows = load_tab(path, "Compliance Matrix")
    grounded = 0
    for row in rows:
        words = content_words(row.get("Requirement Text"))
        pages = cited_pages(row.get("Source Location"), len(page_words))
        near = set().union(*(page_words[p - 1] for p in pages))
        grounded += bool(words) and sum(w in near for w in words) / len(words) >= FOUND_ON_PAGE
    found = [set(content_words(r.get("Requirement Text"))) for r in rows]
    wanted = [w for r in load_tab(reference, "Compliance Matrix") if (w := content_words(r.get("Requirement Text")))]
    matched = sum(any(sum(x in f for x in w) / len(w) >= MATCHES_REFERENCE for f in found) for w in wanted)
    reference_dates = dates(load_tab(reference, "Key Dates"))
    n = max(len(rows), 1)
    return {
        "requirements": len(rows),
        "grounded": grounded / n,
        "ungrounded": len(rows) - grounded,
        "requirement_recall": matched / max(len(wanted), 1),
        "date_recall": len(dates(load_tab(path, "Key Dates")) & reference_dates) / max(len(reference_dates), 1),
        "evaluation_link": sum(bool(r.get("Evaluation Link")) for r in rows) / n,
        "related_ids": sum(bool(r.get("Related IDs")) for r in rows) / n,
    }


def main(pdf: str, reference: str, *candidates: str) -> None:
    result = extract_document(Path(pdf).read_bytes(), Path(pdf).name)
    print(f"reference: {Path(reference).name}")
    print(
        f"{'workbook':42} {'reqs':>5} {'grounded':>9} {'ungrounded':>10} {'req recall':>10} {'date recall':>11}"
        f" {'eval link':>9} {'related':>8}"
    )
    for path in (reference, *candidates):
        s = score(result, Path(reference), Path(path))
        print(
            f"{Path(path).name:42} {s['requirements']:>5} {s['grounded']:>9.0%} {s['ungrounded']:>10}"
            f" {s['requirement_recall']:>10.0%} {s['date_recall']:>11.0%} {s['evaluation_link']:>9.0%}"
            f" {s['related_ids']:>8.0%}"
        )


if __name__ == "__main__":
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    main(*sys.argv[1:])
